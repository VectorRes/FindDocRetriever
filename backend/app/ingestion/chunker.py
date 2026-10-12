"""Turns parsed documents into retrievable chunks with citation metadata."""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from app.ingestion.excel_parser import ParsedSheet, ParsedWorkbook
from app.ingestion.pdf_parser import ParsedPdfDocument, ParsedPdfPage, ParsedPdfTable, ParsedPdfReference


@dataclass
class Chunk:
    sheet_name: str | None
    cell_range: str | None
    text: str
    page_number: int | None = None
    content_type: str = "text"
    table_index: int | None = None
    reference_number: str | None = None
    reference_type: str | None = None
    # How the underlying page text was obtained ("native" or "ocr") and how
    # reliable it is. OCR-derived chunks carry a confidence lower than native
    # text so callers can flag them for verification against the scan.
    extraction_method: str = "native"
    confidence: float | None = None
    confidence_label: str | None = None


# Title/notes rows above the header that are repeated into every row chunk.
MAX_PREAMBLE_ROWS = 3
MAX_PREAMBLE_CHARS = 240


def _header_row(sheet: ParsedSheet) -> tuple[int | None, dict[int, str], list[str]]:
    """The header row, its labels, and the preamble above it.

    Financial statements usually open with a title and notes ("Andina S.A.S. —
    Estado de resultados — Año 2024", "Cifras en pesos colombianos (COP)")
    before the real header ("Concepto | 2024"). So the header is the first row
    with at least two distinct values — a title, even merged across columns,
    has one — falling back to the first row (e.g. single-column sheets). Rows
    above it are kept as preamble: they carry the entity, period and currency
    that every figure in the sheet depends on."""
    if not sheet.cells:
        return None, {}, []
    rows: dict[int, list] = {}
    for cell in sheet.cells:
        if cell.value:
            rows.setdefault(cell.row, []).append(cell)
    if not rows:
        return None, {}, []
    ordered = sorted(rows)
    header_row = next((r for r in ordered if len({c.value for c in rows[r]}) >= 2), ordered[0])
    headers = {cell.column: cell.value for cell in rows[header_row]}
    preamble = [
        " ".join(dict.fromkeys(str(c.value) for c in sorted(rows[r], key=lambda c: c.column)))
        for r in ordered if r < header_row
    ][:MAX_PREAMBLE_ROWS]
    return header_row, headers, preamble


def chunk_sheet(sheet: ParsedSheet) -> list[Chunk]:
    if sheet.is_empty:
        return []
    header_row_number, headers, preamble = _header_row(sheet)
    context = " · ".join(preamble)[:MAX_PREAMBLE_CHARS]
    sheet_label = f"Sheet '{sheet.name}' ({context})" if context else f"Sheet '{sheet.name}'"
    rows: dict[int, list] = {}
    for cell in sheet.cells:
        rows.setdefault(cell.row, []).append(cell)
    chunks: list[Chunk] = []
    for row_number, row_cells in sorted(rows.items()):
        if header_row_number is not None and row_number < header_row_number and preamble:
            continue  # folded into every chunk as context
        if row_number == header_row_number and headers:
            continue
        row_cells = sorted(row_cells, key=lambda c: c.column)
        parts = []
        for cell in row_cells:
            col_label = headers.get(cell.column) or get_column_letter(cell.column)
            piece = f"{col_label}={cell.value if cell.value is not None else '(empty)'}"
            if cell.formula:
                piece += f" (formula: {cell.formula})"
            parts.append(piece)
        row_label = row_cells[0].value or f"row {row_number}"
        text = f"{sheet_label}, {row_label}: " + "; ".join(parts)
        min_col = get_column_letter(row_cells[0].column)
        max_col = get_column_letter(row_cells[-1].column)
        chunks.append(Chunk(sheet_name=sheet.name, cell_range=f"{min_col}{row_number}:{max_col}{row_number}", text=text))
    if headers:
        chunks.append(Chunk(sheet_name=sheet.name, cell_range=None, text=f"{sheet_label} has columns: {', '.join(map(str, headers.values()))}."))
    return chunks


def chunk_workbook(workbook: ParsedWorkbook) -> list[Chunk]:
    return [chunk for sheet in workbook.sheets for chunk in chunk_sheet(sheet)]


def _split_into_windows(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    words = text.split()
    if not words:
        return []
    windows: list[str] = []
    start = 0
    while start < len(words):
        window_words: list[str] = []
        length = 0
        end = start
        while end < len(words):
            projected = length + (1 if window_words else 0) + len(words[end])
            if window_words and projected > chunk_size:
                break
            window_words.append(words[end])
            length = projected
            end += 1
        windows.append(" ".join(window_words))
        if end >= len(words):
            break
        overlap_words = 0
        overlap_len = 0
        i = end - 1
        while i >= start and overlap_len < chunk_overlap:
            overlap_len += len(words[i]) + 1
            overlap_words += 1
            i -= 1
        start = max(end - overlap_words, start + 1)
    return windows


def chunk_pdf_page(page: ParsedPdfPage, chunk_size: int = 1000, chunk_overlap: int = 150) -> list[Chunk]:
    return [
        Chunk(
            None,
            None,
            window,
            page.number,
            extraction_method=page.extraction_method,
            confidence=page.confidence,
            confidence_label=page.confidence_label,
        )
        for window in _split_into_windows(page.text, chunk_size, chunk_overlap)
    ]


def _table_row_chunks(table: ParsedPdfTable) -> list[Chunk]:
    chunks: list[Chunk] = []
    for row in table.rows:
        padded = row + [""] * max(0, len(table.headers) - len(row))
        label = padded[0] or "(unlabelled row)"
        values = [f"{header}={value}" for header, value in zip(table.headers[1:], padded[1:]) if value]
        text = f"Financial table page {table.page_number}, row '{label}': " + "; ".join(values)
        chunks.append(
            Chunk(None, None, text, table.page_number, "table_row", table.table_index)
        )
    return chunks


def _reference_chunks(reference: ParsedPdfReference) -> list[Chunk]:
    text = f"{reference.reference_type.title()} {reference.reference_number}"
    if reference.title:
        text += f": {reference.title}"
    text += f" (page {reference.page_number})"
    return [Chunk(None, None, text + f". Source: {reference.text}", reference.page_number, "reference", None, reference.reference_number, reference.reference_type)]


def chunk_pdf(document: ParsedPdfDocument, chunk_size: int = 1000, chunk_overlap: int = 150) -> list[Chunk]:
    chunks: list[Chunk] = []
    for page in document.pages:
        chunks.extend(chunk_pdf_page(page, chunk_size=chunk_size, chunk_overlap=chunk_overlap))
    for table in document.tables:
        chunks.extend(_table_row_chunks(table))
    for reference in document.references:
        chunks.extend(_reference_chunks(reference))
    return chunks

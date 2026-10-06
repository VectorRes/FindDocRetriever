"""Excel export of a comparison table (ID-HU-FE-003), citations included.

Built from the comparison result the analyst is looking at (posted back by
the frontend), not recomputed: re-running extraction could return a
different figure than the one on screen.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas import ComparisonOut, CitationOut

_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_MISMATCH_FILL = PatternFill("solid", fgColor="FDE2E1")
_NUMBER_FORMAT = "#,##0.00"


def _source(citation: CitationOut | None) -> str:
    if citation is None:
        return ""
    parts = [citation.document_filename]
    if citation.sheet_name:
        parts.append(f"sheet '{citation.sheet_name}'")
    if citation.cell_range:
        parts.append(f"cells {citation.cell_range}")
    if citation.page_number is not None:
        parts.append(f"page {citation.page_number}")
    return ", ".join(parts)


def comparison_workbook(comparison: ComparisonOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    ws.append([f"Comparison: {comparison.metric}"])
    ws["A1"].font = Font(bold=True, size=13)
    if comparison.period:
        ws.append(["Period", comparison.period])
    if comparison.across_periods:
        status = "Across periods (variance only)"
    else:
        status = {True: "Reconciles", False: "Does NOT reconcile", None: "Can't determine"}[comparison.reconciles]
    ws.append(["Result", status])
    if comparison.comparison_currency:
        ws.append(["Compared in", comparison.comparison_currency])
    ws.append([])

    headers = [
        "Document", "Version", "Baseline", "Period", "Line item", "Value as stated", "Currency",
        "Value (base units)", f"Compared value ({comparison.comparison_currency or '-'})",
        "Variance vs baseline", "Variance %", "Matches baseline", "Source", "Confidence", "Notes",
    ]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL

    for row in comparison.rows:
        ws.append([
            row.filename,
            f"v{row.version_number}{'' if row.is_current else ' (superseded)'}",
            "yes" if row.is_baseline else "",
            row.period,
            row.label,
            row.value_text if row.found else "not found",
            row.currency,
            row.value,
            row.comparable_value,
            row.variance_abs,
            None if row.variance_pct is None else row.variance_pct / 100,
            {True: "yes", False: "NO", None: ""}[row.matches_baseline],
            _source(row.citation),
            row.confidence if row.found else "",
            row.reason or "",
        ])
        current = ws.max_row
        for column in (8, 9, 10):
            ws.cell(current, column).number_format = _NUMBER_FORMAT
        ws.cell(current, 11).number_format = "0.00%"
        if row.matches_baseline is False:
            for cell in ws[current]:
                cell.fill = _MISMATCH_FILL

    ws.append([])
    for rate in comparison.exchange_rates:
        ws.append([
            "Exchange rate",
            f"1 {rate.from_currency} = {rate.rate:,.6g} {rate.to_currency}",
            f"as written: {rate.rate_text}",
            _source(rate.citation),
        ])
    for note in comparison.notes:
        ws.append(["Note", note])
    if comparison.escalation:
        ws.append(["Suggested escalation", comparison.escalation.team, comparison.escalation.reason])

    widths = [34, 14, 9, 12, 26, 16, 9, 18, 20, 18, 11, 10, 48, 11, 50]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(header_row, index).column_letter].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (13, 15))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

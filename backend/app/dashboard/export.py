"""Excel export of the ratios dashboard (ID-HU-FE-004), sources included.

Built from the dashboard the analyst is looking at (posted back by the
frontend), like the comparison export: three sheets — the ratios by period,
every input with its citation, and the outliers and notes.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.figures.export import _source
from app.schemas import DashboardOut, DashboardPointOut

_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_OUTLIER_FILL = PatternFill("solid", fgColor="FEF3C7")
_STATUS_TEXT = {
    "needs_review": "held for review",
    "not_calculable": "not calculable",
    "not_calculated": "not calculated",
    "restricted": "restricted",
}


def _text(value):
    # A label from a document starting with "=" must not become a formula.
    return f"'{value}" if isinstance(value, str) and value.startswith(("=", "+", "@")) else value


def _shown(point: DashboardPointOut):
    if point.status != "calculated" or point.value is None:
        return _STATUS_TEXT.get(point.status, point.status)
    return point.value


def _header(ws, row: list[str]) -> int:
    ws.append(row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    return ws.max_row


def dashboard_workbook(dashboard: DashboardOut) -> bytes:
    wb = Workbook()
    periods = [p.period for p in dashboard.periods]

    ws = wb.active
    ws.title = "Ratios"
    ws.append([f"Key financial ratios: {dashboard.entity}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    header_row = _header(ws, ["Category", "Ratio", "Formula", *periods])
    for ratio in dashboard.ratios:
        ws.append([ratio.category, ratio.name, ratio.formula, *(_shown(p) for p in ratio.points)])
        row = ws.max_row
        for index, point in enumerate(ratio.points, start=4):
            cell = ws.cell(row, index)
            if isinstance(cell.value, float):
                cell.number_format = "0.0%" if ratio.unit == "percent" else '0.00"x"'
            if point.outlier is not None:
                cell.fill = _OUTLIER_FILL
    for letter, width in zip("ABC", (14, 26, 40)):
        ws.column_dimensions[letter].width = width
    for index in range(4, 4 + len(periods)):
        ws.column_dimensions[ws.cell(header_row, index).column_letter].width = 16

    sources = wb.create_sheet("Inputs and sources")
    _header(sources, ["Ratio", "Period", "Input", "Line item in source", "Value as stated", "Currency",
                      "Document", "Source", "Status", "Notes"])
    for ratio in dashboard.ratios:
        for point in ratio.points:
            for item in point.inputs:
                sources.append([_text(v) for v in (
                    ratio.name, point.period, item.name, item.label, item.value_text if item.found else "not found",
                    item.currency, item.document_filename, _source(item.citation),
                    _STATUS_TEXT.get(point.status, point.status), item.reason or "",
                )])
            if not point.inputs and point.reasons:
                sources.append([ratio.name, point.period, "", "", "", "", "", "",
                                _STATUS_TEXT.get(point.status, point.status), " ".join(point.reasons)])
    for letter, width in zip("ABCDEFGHIJ", (26, 10, 22, 30, 16, 9, 32, 44, 16, 60)):
        sources.column_dimensions[letter].width = width
    for row in sources.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (8, 10))

    review = wb.create_sheet("Outliers and notes")
    _header(review, ["Ratio", "Period", "Compared with", "Change", "Review"])
    for ratio in dashboard.ratios:
        for point in ratio.points:
            if point.outlier is not None:
                outlier = point.outlier
                review.append([
                    ratio.name, point.period, outlier.previous_period or "",
                    outlier.change if outlier.change is not None else "", _text(outlier.message),
                ])
                if outlier.change is not None:
                    review.cell(review.max_row, 4).number_format = "+0%;-0%"
    review.append([])
    for note in dashboard.notes:
        review.append(["Note", _text(note)])
    for period in dashboard.periods:
        for note in period.notes:
            review.append([f"Note ({period.period})", _text(note)])
    for letter, width in zip("ABCDE", (26, 12, 14, 10, 90)):
        review.column_dimensions[letter].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

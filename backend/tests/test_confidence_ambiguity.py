"""ID-HU-BE-009 — low-confidence answers, ambiguous questions and escalation.

Like test_qa_generation.py, these never call the real OpenAI API: the chat
model is monkeypatched with a fake that answers each structured-output call
with the response scripted for its schema — the triage pre-check
(LLMTriage) or answer generation (LLMGroundedAnswer) — so the graph's
enforcement logic is tested deterministically.
"""
from uuid import UUID

import openpyxl
import pytest

import app.qa.graph as qa_graph_module
from app.escalation import classify_topic, resolve_topic
from app.ingestion.service import ingest_excel_file
from app.qa.confidence import assess_confidence, unverified_figures
from app.qa.conversation import get_recent_history
from app.qa.graph import clarification_for, get_qa_graph
from app.qa.schemas import (
    GroundedStatement,
    LLMCitation,
    LLMGroundedAnswer,
    LLMStatement,
    LLMTriage,
    ResolvedCitation,
)
from app.qa.service import answer_question
from app.retrieval.service import RetrievedChunk


def make_chunk(text: str, **overrides) -> RetrievedChunk:
    defaults = dict(
        document_id="doc-1",
        document_filename="income_statement.xlsx",
        sheet_name="P&L",
        cell_range="A2:B2",
        page_number=None,
        text=text,
    )
    defaults.update(overrides)
    return RetrievedChunk(**defaults)


class _FakeChatModel:
    """Answers each structured-output call with the response scripted for its
    schema. A scripted response that is an exception is raised instead, and a
    schema with no scripted response fails the test (an unexpected call)."""

    def __init__(self, responses: dict):
        self._responses = responses

    def with_structured_output(self, schema):
        if schema not in self._responses:
            raise AssertionError(f"unexpected model call for {schema.__name__}")
        response = self._responses[schema]

        class _Structured:
            def invoke(self, messages):
                if isinstance(response, Exception):
                    raise response
                return response

        return _Structured()


def _patch(monkeypatch, raw=None, triage=None):
    responses = {LLMTriage: triage if triage is not None else LLMTriage()}
    if raw is not None:
        responses[LLMGroundedAnswer] = raw
    monkeypatch.setattr(qa_graph_module, "get_chat_model", lambda: _FakeChatModel(responses))


def _run(monkeypatch, question, chunks, raw=None, notice=None, triage=None):
    _patch(monkeypatch, raw=raw, triage=triage)
    return get_qa_graph().invoke(
        {"question": question, "chunks": chunks, "history": [], "restriction_notice": notice}
    )["answer"]


def _statement(text: str, source_id: str = "S1") -> LLMStatement:
    return LLMStatement(text=text, citations=[LLMCitation(source_id=source_id)])


def _grounded(text: str, source_text: str) -> GroundedStatement:
    citation = ResolvedCitation(
        source_id="S1", document_id="d", document_filename="f.xlsx", sheet_name=None,
        cell_range=None, page_number=None, reference_number=None, text=source_text,
    )
    return GroundedStatement(text=text, citations=[citation])


# --- figure verification -------------------------------------------------

@pytest.mark.parametrize(
    ("statement", "source"),
    [
        ("Net income was 81,000.", "Net income: 81000"),
        ("La utilidad neta fue de 1.234.567,89.", "Utilidad neta=1234567.89"),
        ("Revenue was 3.4 million.", "Revenue: 3,401,234"),
        ("Revenue was 3,401,234.", "Revenue (in thousands): 3,401.234"),
        ("Margin was 12.5% in 2024 Q3.", "Margin 12.5"),
    ],
)
def test_figures_present_in_the_source_are_verified(statement, source):
    assert unverified_figures(_grounded(statement, source)) == []


def test_a_figure_missing_from_the_source_is_flagged():
    assert unverified_figures(_grounded("Revenue was 95,000.", "Revenue: 81,000")) == ["95,000"]


def test_years_and_small_labels_are_not_treated_as_figures():
    assert unverified_figures(_grounded("In 2024 Q3, see note 12.", "Note twelve")) == []


# --- "no confident answer" -------------------------------------------------

def test_statement_with_fabricated_figure_is_removed_not_shown(monkeypatch):
    raw = LLMGroundedAnswer(statements=[
        _statement("Net income was 120,000."),
        _statement("Supplier concentration risk was 35.7%."),  # not in any source
    ])
    answer = _run(monkeypatch, "Net income and supplier risk?", [make_chunk("Net income: 120,000")], raw)

    assert [s.text for s in answer.statements] == ["Net income was 120,000."]
    assert answer.confidence.level != "high"
    assert any("35.7" in reason for reason in answer.confidence.reasons)


def test_no_verifiable_statement_returns_no_confident_answer_with_escalation(monkeypatch):
    raw = LLMGroundedAnswer(statements=[_statement("Supplier concentration risk was 35.7%.")])
    answer = _run(monkeypatch, "What is our supplier concentration risk?", [make_chunk("Revenue: 81,000")], raw)

    assert answer.grounded is False
    assert answer.statements == []
    assert answer.confidence is None
    assert answer.escalation is not None


def test_no_sources_returns_no_confident_answer_without_calling_the_model(monkeypatch):
    # Only the (skipped) triage is scripted: a generation call would fail the test.
    _patch(monkeypatch)
    answer = get_qa_graph().invoke({"question": "What is our tax exposure?", "chunks": [], "history": []})["answer"]

    assert answer.grounded is False
    assert answer.escalation.team == "Tax"


def test_no_answer_with_restricted_matches_still_suggests_a_team(monkeypatch):
    raw = LLMGroundedAnswer(statements=[], unsupported=True)
    answer = _run(monkeypatch, "Executive bonus?", [make_chunk("Public data")], raw, notice="Part is restricted.")

    assert answer.escalation.team == "HR/Compensation"
    assert "sources you can access" in answer.escalation.reason


# --- confidence ------------------------------------------------------------

def test_explicit_verified_answer_is_high_confidence(monkeypatch):
    raw = LLMGroundedAnswer(statements=[_statement("Net income was 120,000.")], self_confidence="high")
    answer = _run(monkeypatch, "Net income?", [make_chunk("Net income: 120,000")], raw)

    assert answer.confidence.level == "high"
    assert answer.confidence.reasons == []
    assert answer.escalation is None


def test_judgment_question_is_low_confidence_and_suggests_accounting(monkeypatch):
    """Even when the generation model claims high confidence, a question that
    asks for an accounting treatment is never presented as a confident answer."""
    raw = LLMGroundedAnswer(
        statements=[_statement("The contract includes a 12-month service component billed upfront.")],
        self_confidence="high",
    )
    answer = _run(
        monkeypatch,
        "How should we recognise revenue for the Acme service contract?",
        [make_chunk("Acme contract: 12-month service component, billed upfront", sheet_name=None)],
        raw,
        triage=LLMTriage(requires_judgment=True, topic="revenue_recognition"),
    )

    assert answer.confidence.level == "low"
    assert any("professional judgment" in r for r in answer.confidence.reasons)
    assert answer.escalation.team == "Accounting/Consolidation"


def test_model_reported_low_confidence_alone_marks_the_answer_low(monkeypatch):
    raw = LLMGroundedAnswer(statements=[_statement("Net income was 120,000.")], self_confidence="low")
    answer = _run(monkeypatch, "Net income?", [make_chunk("Net income: 120,000")], raw)

    assert answer.confidence.level == "low"
    assert answer.escalation is not None


def test_a_failed_triage_does_not_block_the_answer(monkeypatch):
    raw = LLMGroundedAnswer(statements=[_statement("Net income was 120,000.")])
    answer = _run(
        monkeypatch, "Net income?", [make_chunk("Net income: 120,000")], raw,
        triage=RuntimeError("triage model unavailable"),
    )
    assert [s.text for s in answer.statements] == ["Net income was 120,000."]


def test_weakly_related_sources_lower_confidence_when_scores_exist():
    weak = assess_confidence(1, 1, [], "high", cited_scores=[0.2])
    strong = assess_confidence(1, 1, [], "high", cited_scores=[0.9])
    assert weak.score < strong.score
    assert any("weakly related" in r for r in weak.reasons)


def test_dropped_statements_lower_confidence():
    assessment = assess_confidence(
        kept_count=1, proposed_count=3, removed_figures=[], self_confidence="high", cited_scores=[]
    )
    assert assessment.level in {"medium", "low"}
    assert "2 of 3" in assessment.reasons[0]


# --- ambiguity -------------------------------------------------------------

def test_ambiguous_question_asks_for_clarification_without_generating_an_answer(monkeypatch):
    triage = LLMTriage(
        metric_in_sources=True,
        entities_with_metric=["Andina S.A.S.", "Pacífico S.A."],
        periods_with_metric=["2024"],
    )
    chunks = [
        make_chunk("Andina S.A.S. 2024 revenue: 500,000"),
        make_chunk("Pacífico S.A. 2024 revenue: 320,000", document_id="doc-2"),
    ]
    # No LLMGroundedAnswer scripted: the generation model must not be called.
    answer = _run(monkeypatch, "What was revenue?", chunks, triage=triage)

    assert answer.needs_clarification is True
    assert answer.statements == []
    assert answer.clarification_question == "Which entity do you mean?"
    assert answer.clarification_options == ["Andina S.A.S.", "Pacífico S.A."]
    assert answer.confidence is None


def test_without_sources_there_is_nothing_to_clarify(monkeypatch):
    _patch(monkeypatch, triage=AMBIGUOUS_PERIOD)
    answer = get_qa_graph().invoke({"question": "Revenue?", "chunks": [], "history": []})["answer"]

    assert answer.needs_clarification is False
    assert answer.grounded is False


def test_clarification_is_kept_in_history_so_the_reply_resolves(monkeypatch, tmp_path, db_session):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Andina 2024 revenue"
    ws["B1"] = 500000
    path = tmp_path / "andina.xlsx"
    wb.save(path)

    document = ingest_excel_file(db_session, filename="andina_be009.xlsx", file_path=path)
    document.confidentiality_tag = "public"
    for chunk in document.chunks:
        chunk.confidentiality_tag = "public"
    db_session.commit()
    try:
        _patch(monkeypatch, triage=AMBIGUOUS_PERIOD)
        answer = answer_question(db_session, question="What was revenue?")

        history = get_recent_history(db_session, UUID(answer.session_id))
        assert history[-1] == ("What was revenue?", "(asked for clarification) Which period do you mean?")
    finally:
        db_session.delete(document)
        db_session.commit()


AMBIGUOUS_PERIOD = LLMTriage(metric_in_sources=True, periods_with_metric=["Q2 2024", "Q3 2024"])


@pytest.mark.parametrize(
    ("triage", "expected"),
    [
        # Several entities and periods, nothing specified: ask for both.
        (LLMTriage(metric_in_sources=True, entities_with_metric=["Andina", "Pacífico"],
                   periods_with_metric=["2024", "2025"]),
         ("Which entity and period do you mean?",
          ["Andina – 2024", "Andina – 2025", "Pacífico – 2024", "Pacífico – 2025"])),
        # Period given (or a reply to an earlier clarification): only the entity is missing.
        (LLMTriage(metric_in_sources=True, entities_with_metric=["Andina", "Pacífico"],
                   periods_with_metric=["2024", "2025"], question_specifies_period=True),
         ("Which entity do you mean?", ["Andina", "Pacífico"])),
        # Both specified: not ambiguous.
        (LLMTriage(metric_in_sources=True, entities_with_metric=["Andina", "Pacífico"],
                   periods_with_metric=["2024", "2025"], question_specifies_entity=True,
                   question_specifies_period=True), None),
        # One entity, one period: nothing to choose.
        (LLMTriage(metric_in_sources=True, entities_with_metric=["Andina"], periods_with_metric=["2024"]), None),
        # Duplicates differing only in case/spacing count once.
        (LLMTriage(metric_in_sources=True, periods_with_metric=["Q3 2024", "q3  2024"]), None),
        # Metric not in the sources (e.g. "supplier concentration risk"): never ambiguous.
        (LLMTriage(metric_in_sources=False, periods_with_metric=["Q2 2024", "Q3 2024"]), None),
    ],
)
def test_clarification_is_decided_from_triage_facts(triage, expected):
    assert clarification_for(triage) == expected


def test_judgment_questions_tell_generation_to_report_facts_only(monkeypatch):
    prompts = []

    class _Recording(_FakeChatModel):
        def with_structured_output(self, schema):
            inner = super().with_structured_output(schema)

            class _Wrap:
                def invoke(self, messages):
                    prompts.append(messages[-1]["content"])
                    return inner.invoke(messages)

            return _Wrap()

    responses = {
        LLMTriage: LLMTriage(requires_judgment=True),
        LLMGroundedAnswer: LLMGroundedAnswer(statements=[_statement("Billed 100% upfront.")]),
    }
    monkeypatch.setattr(qa_graph_module, "get_chat_model", lambda: _Recording(responses))
    get_qa_graph().invoke({"question": "How should we recognise it?", "chunks": [make_chunk("Billed 100% upfront")], "history": []})

    assert "professional judgment" in prompts[-1]


# --- escalation routing ----------------------------------------------------

@pytest.mark.parametrize(
    ("text", "topic"),
    [
        ("How should we recognize revenue for this contract under IFRS 15?", "revenue_recognition"),
        ("¿Cuál es el tratamiento contable del arrendamiento?", "accounting"),
        ("Does net income reconcile with the cash flow statement?", "consolidation"),
        ("What is the deferred tax balance?", "tax"),
        ("¿Cuál fue el flujo de caja libre?", "treasury"),
        ("What were executive bonuses?", "compensation"),
        ("Are there open litigation contingencies?", "legal"),
        ("¿Cuál es el presupuesto de marketing?", "budget_forecast"),
        ("¿Cuál fue la rentabilidad del trimestre?", "general"),   # "renta" isn't tax
        ("When is the next product release?", "general"),           # "lease" isn't accounting
    ],
)
def test_classify_topic(text, topic):
    assert classify_topic(text) == topic


def test_model_topic_wins_over_keywords_unless_generic():
    assert resolve_topic("tax", "revenue recognition question") == "tax"
    assert resolve_topic("general", "revenue recognition question") == "revenue_recognition"
    assert resolve_topic("not-a-topic", "payroll costs") == "compensation"

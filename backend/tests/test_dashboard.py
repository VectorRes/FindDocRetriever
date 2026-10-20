"""ID-HU-FE-004 — key financial figures and ratios dashboard.

Catalog, period and outlier logic run without a database. The dashboard
tests ingest real workbooks (fake embeddings, see conftest.py) so the catalog
reads real chunks, and script the figure extraction (no model): each test
uses its own company name, so other tests' documents don't interfere.
"""
import io
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.dashboard.catalog import entity_from_filename, parse_period, period_of_column
from app.dashboard.outliers import Outlier, detect_outliers, explain
from app.dashboard.service import DashboardError, get_dashboard, refresh
from app.db.models import DashboardRun, RetrievalAuditLog
from app.db.session import SessionLocal
from app.figures.extraction import ExtractedFigure, ExtractionServiceError
from app.ingestion.service import ingest_excel_file
from app.qa.schemas import ResolvedCitation
from app.ratios.formulas import LINE_ITEMS

# --- catalog: entities and periods ---------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "entity"),
    [
        ("EERR_Andina_2024.xlsx", "Andina"),
        ("Balance_Andina_2024_v2.xlsx", "Andina"),
        ("EEFF EL ECLIPSE S.A.S DICIEMBRE 2024.xlsx", "Eclipse S.A.S"),
        ("Alphabet_Q2_2026_earnings_release.pdf", "Alphabet"),
        ("Pacifico - Estado de situacion financiera 2025.xlsx", "Pacifico"),
    ],
)
def test_entity_comes_from_the_filename(filename, entity):
    assert entity_from_filename(filename) == entity


@pytest.mark.parametrize(
    ("column", "period"),
    [
        ("2024", "2024"),
        ("Año 2024", "2024"),
        ("Q1 2024", "Q1 2024"),
        ("Three Months Ended June 30, 2026", "Q2 2026"),
        ("Tres meses terminados el 30 de junio 2026", "Q2 2026"),
        ("June 27, 2026", "Q2 2026"),
        ("Al 31 de diciembre de 2025", "2025"),
        ("Years Ended September 27, 2025", "2025"),
        ("Nine Months Ended June 27, 2026", None),
        ("Year to date 2026", None),
        ("Presupuesto 2025", None),
        ("Variación 2025", None),
        ("Concepto", None),
    ],
)
def test_period_of_a_column_heading(column, period):
    found = period_of_column(column)
    assert (found.label if found else None) == period


@pytest.mark.parametrize(("text", "label"), [("FY2024", "2024"), ("2024", "2024"), ("Q3 2026", "Q3 2026"), ("3Q26", "Q3 2026")])
def test_typed_periods(text, label):
    assert parse_period(text).label == label


def test_periods_sort_chronologically():
    periods = [parse_period(p) for p in ("2025", "Q3 2024", "2024", "Q1 2025")]
    assert [p.label for p in sorted(periods)] == ["Q3 2024", "2024", "Q1 2025", "2025"]


# --- outliers -----------------------------------------------------------------------------------


def test_roe_jump_from_12_to_34_percent_is_flagged():
    flagged = detect_outliers([("Q1", 0.12), ("Q2", 0.34)], "percent")
    assert list(flagged) == ["Q2"]
    assert flagged["Q2"].previous_period == "Q1"
    assert flagged["Q2"].change == pytest.approx(0.22 / 0.12)


@pytest.mark.parametrize(
    ("points", "unit"),
    [
        ([("2023", 0.02), ("2024", 0.035)], "percent"),  # +75% but only 1.5 pp
        ([("2023", 1.0), ("2024", 1.4)], "times"),  # +0.4x, +40%
        ([("2023", 0.12), ("2024", 0.13), ("2025", 0.125)], "percent"),
    ],
)
def test_small_moves_are_not_flagged(points, unit):
    assert detect_outliers(points, unit) == {}


def test_a_value_far_from_a_long_series_is_flagged():
    points = [("2021", 2.0), ("2022", 2.1), ("2023", 2.05), ("2024", 2.1), ("2025", 3.2)]
    assert "2025" in detect_outliers(points, "times")


def test_roe_outlier_is_explained_by_the_leverage_change():
    outlier = Outlier("2024", 0.13, (0.34 - 0.13) / 0.13, 0.21, "period_change")
    explained = explain(
        "roe", "Return on equity (ROE)", outlier, "2025",
        {"roe": 0.13, "roa": 0.078, "net_margin": 0.12, "asset_turnover": 0.65},
        {"roe": 0.34, "roa": 0.081, "net_margin": 0.121, "asset_turnover": 0.67},
        [], [], "percent",
    )
    names = [d.name for d in explained.drivers]
    assert names == ["Net margin", "Asset turnover", "Equity multiplier"]
    assert "equity multiplier" in explained.message and "leverage" in explained.message
    assert "Review the underlying drivers" in explained.message


def test_other_ratios_are_explained_by_their_inputs():
    outlier = Outlier("2024", 0.66, 3.8, 2.5, "period_change")
    explained = explain(
        "debt_to_equity", "Debt-to-equity", outlier, "2025", {}, {},
        [{"key": "total_equity", "name": "Total equity", "found": True, "value": 5_000_000},
         {"key": "total_liabilities", "name": "Total liabilities", "found": True, "value": 3_300_000}],
        [{"key": "total_equity", "name": "Total equity", "found": True, "value": 2_060_000},
         {"key": "total_liabilities", "name": "Total liabilities", "found": True, "value": 6_540_000}],
        "times",
    )
    assert "total equity" in explained.message
    assert {d.unit for d in explained.drivers} == {"amount"}


# --- the dashboard ----------------------------------------------------------------------------

INCOME = {  # year: line items (label, value)
    2023: {"revenue": ("Ingresos", 5_000_000), "ebitda": ("EBITDA", 1_400_000), "net_income": ("Utilidad neta", 600_000),
           "operating_income": ("Utilidad operacional", 1_150_000), "interest_expense": ("Gastos financieros", 180_000)},
    2024: {"revenue": ("Ingresos", 5_400_000), "ebitda": ("EBITDA", 1_500_000), "net_income": ("Utilidad neta", 650_000),
           "operating_income": ("Utilidad operacional", 1_240_000), "interest_expense": ("Gastos financieros", 190_000)},
    2025: {"revenue": ("Ingresos", 5_800_000), "ebitda": ("EBITDA", 1_600_000), "net_income": ("Utilidad neta", 700_400),
           "operating_income": ("Utilidad operacional", 1_330_000), "interest_expense": ("Gastos financieros", 260_000)},
}
BALANCE = {
    2023: {"total_assets": ("Total activos", 8_000_000), "total_equity": ("Total patrimonio", 5_000_000),
           "total_liabilities": ("Total pasivos", 3_000_000), "current_assets": ("Total activo corriente", 3_000_000),
           "current_liabilities": ("Total pasivo corriente", 1_500_000)},
    2024: {"total_assets": ("Total activos", 8_300_000), "total_equity": ("Total patrimonio", 5_000_000),
           "total_liabilities": ("Total pasivos", 3_300_000), "current_assets": ("Total activo corriente", 3_100_000),
           "current_liabilities": ("Total pasivo corriente", 1_550_000)},
    2025: {"total_assets": ("Total activos", 8_600_000), "total_equity": ("Total patrimonio", 2_060_000),
           "total_liabilities": ("Total pasivos", 6_540_000), "current_assets": ("Total activo corriente", 3_050_000),
           "current_liabilities": ("Total pasivo corriente", 2_100_000)},
}


class StatementExtractor:
    """Plays extract_figure from what each ingested workbook holds."""

    def __init__(self):
        self.tables: dict[str, dict] = {}  # document id -> {key: (label, value, sheet)}
        self.filenames: dict[str, str] = {}
        self.calls: list[tuple[str, str, str | None]] = []
        self.fail = False

    def __call__(self, db, document_id, metric, user_roles=None, user_id=None, period=None, **_):
        key = next(k for k, item in LINE_ITEMS.items() if item.query == metric)
        self.calls.append((str(document_id), key, period))
        if self.fail:
            return ExtractedFigure(document_id=str(document_id), found=False, reason="down", error=True)
        found = self.tables.get(str(document_id), {}).get(key)
        if found is None:
            return ExtractedFigure(document_id=str(document_id), found=False, reason=f"No {metric}.")
        label, value, sheet = found
        citation = ResolvedCitation(
            source_id="S1", document_id=str(document_id), document_filename=self.filenames[str(document_id)],
            sheet_name=sheet, cell_range="A4:B4", page_number=None, reference_number=None, text=f"{label}: {value}",
        )
        return ExtractedFigure(
            document_id=str(document_id), found=True, value=float(value), value_text=f"{value:,}", currency="COP",
            period=period, label=label, citation=citation, confidence="high",
        )


@pytest.fixture
def company(db_session, tmp_path):
    """Ingests statements for a company named uniquely for the test."""
    name = f"Zeta{uuid.uuid4().hex[:6]}"
    extractor = StatementExtractor()
    created = []

    def add(kind: str, year: int, items: dict, tag: str = "internal", name_override: str | None = None):
        sheet = "Estado de resultados" if kind == "EERR" else "Balance general"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append([f"{name} — {sheet} — Año {year}"])
        ws.append(["Concepto", str(year)])
        for label, value in items.values():
            ws.append([label, value])
        path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
        wb.save(path)
        filename = name_override or f"{kind}_{name}_{year}.xlsx"
        document = ingest_excel_file(db_session, filename=filename, file_path=path)
        document.confidentiality_tag = tag
        for chunk in document.chunks:
            chunk.confidentiality_tag = tag
        db_session.commit()
        created.append(document)
        extractor.tables[str(document.id)] = {k: (label, value, sheet) for k, (label, value) in items.items()}
        extractor.filenames[str(document.id)] = filename
        return document

    add.name = name
    add.extractor = extractor
    yield add
    for document in created:
        db_session.delete(document)
    session = SessionLocal()
    session.query(DashboardRun).filter(DashboardRun.entity == name).delete()
    session.commit()
    session.close()
    db_session.commit()


def _ratio(dashboard, key):
    return next(r for r in dashboard.ratios if r.key == key)


def _point(dashboard, key, period):
    return next(p for p in _ratio(dashboard, key).points if p.period == period)


def test_main_ebitda_margin_with_citations_to_revenue_and_ebitda(db_session, company):
    company("EERR", 2024, INCOME[2024])
    company("Balance", 2024, BALANCE[2024])

    dashboard = refresh(db_session, company.name, extractor=company.extractor)

    point = _point(dashboard, "ebitda_margin", "2024")
    assert point.status == "calculated"
    assert point.value == pytest.approx(1_500_000 / 5_400_000)
    cited = {c.text.split(":")[0] for c in point.citations}
    assert cited == {"EBITDA", "Ingresos"}
    assert [p.period for p in dashboard.periods] == ["2024"]


def test_trend_over_periods_matches_hand_calculation(db_session, company):
    for year in (2023, 2024, 2025):
        company("EERR", year, INCOME[year])
        company("Balance", year, BALANCE[year])

    dashboard = refresh(db_session, company.name, extractor=company.extractor)

    roe = _ratio(dashboard, "roe")
    assert [p.period for p in roe.points] == ["2023", "2024", "2025"]
    assert [p.value for p in roe.points] == pytest.approx([0.12, 0.13, 700_400 / 2_060_000])
    # Each period only reads that period's documents.
    years = {(doc, period) for doc, _, period in company.extractor.calls}
    for doc, period in years:
        assert company.extractor.filenames[doc].endswith(f"_{period}.xlsx")


def test_alternate_2_outlier_flagged_with_leverage_driver(db_session, company):
    for year in (2023, 2024, 2025):
        company("EERR", year, INCOME[year])
        company("Balance", year, BALANCE[year])

    dashboard = refresh(db_session, company.name, extractor=company.extractor)

    roe = _ratio(dashboard, "roe")
    assert [p.period for p in roe.points if p.outlier] == ["2025"]
    outlier = _point(dashboard, "roe", "2025").outlier
    assert outlier.previous_period == "2024"
    assert [d.name for d in outlier.drivers][-1] == "Equity multiplier"
    assert "leverage" in outlier.message
    assert _point(dashboard, "net_margin", "2025").outlier is None


def test_alternate_1_missing_balance_sheet_is_insufficient_data(db_session, company):
    company("EERR", 2024, INCOME[2024])

    dashboard = refresh(db_session, company.name, extractor=company.extractor)

    for key in ("debt_to_equity", "debt_ratio", "current_ratio", "roe", "roa"):
        point = _point(dashboard, key, "2024")
        assert point.status == "unavailable" and point.value is None
        assert point.missing_statements == ["balance_sheet"]
        assert point.reasons == [f"Insufficient data: the balance sheet for 2024 hasn't been ingested for {company.name}."]
    assert _point(dashboard, "interest_coverage", "2024").status == "calculated"
    # Balance-sheet inputs were never sent to extraction.
    assert not {key for _, key, _ in company.extractor.calls} & {"total_assets", "total_equity", "current_assets"}


def test_criterion_3_ratio_lists_every_contributing_document(db_session, company):
    income = company("EERR", 2024, INCOME[2024])
    balance = company("Balance", 2024, BALANCE[2024])

    dashboard = refresh(db_session, company.name, extractor=company.extractor)

    point = _point(dashboard, "roa", "2024")
    assert point.cross_document is True
    assert {c.document_id for c in point.citations} == {str(income.id), str(balance.id)}


def test_opening_the_dashboard_reads_stored_results_without_extraction(db_session, company):
    company("EERR", 2024, INCOME[2024])
    company("Balance", 2024, BALANCE[2024])
    refresh(db_session, company.name, extractor=company.extractor)
    calls = len(company.extractor.calls)

    dashboard = get_dashboard(db_session, company.name)

    assert len(company.extractor.calls) == calls
    assert _point(dashboard, "roe", "2024").status == "calculated"
    assert dashboard.periods[0].calculated and not dashboard.periods[0].stale


def test_new_document_makes_the_period_stale_until_refreshed(db_session, company):
    company("EERR", 2024, INCOME[2024])
    refresh(db_session, company.name, extractor=company.extractor)

    company("Balance", 2024, BALANCE[2024])
    stale = get_dashboard(db_session, company.name)
    assert stale.periods[0].stale and _point(stale, "roe", "2024").stale
    assert any("refresh" in note for note in stale.notes)

    fresh = refresh(db_session, company.name, extractor=company.extractor)
    assert not fresh.periods[0].stale
    assert _point(fresh, "roe", "2024").status == "calculated"


def test_uncalculated_periods_show_as_not_calculated(db_session, company):
    company("EERR", 2023, INCOME[2023])
    company("EERR", 2024, INCOME[2024])
    refresh(db_session, company.name, periods=["2024"], extractor=company.extractor)

    dashboard = get_dashboard(db_session, company.name)

    assert [(p.period, p.calculated) for p in dashboard.periods] == [("2023", False), ("2024", True)]
    assert _point(dashboard, "net_margin", "2023").status == "not_calculated"


def test_typed_period_without_documents_is_insufficient_data(db_session, company):
    company("EERR", 2024, INCOME[2024])

    dashboard = refresh(db_session, company.name, periods=["2022"], extractor=company.extractor)

    point = _point(dashboard, "net_margin", "2022")
    assert point.status == "unavailable"
    assert point.missing_statements == ["income_statement"]


def test_viewer_without_access_sees_no_values_and_is_logged(db_session, company):
    company("EERR", 2024, INCOME[2024], tag="restricted:compensation")
    company("Balance", 2024, BALANCE[2024])
    refresh(db_session, company.name, user_roles="admin", extractor=company.extractor)
    viewer = f"analyst-{uuid.uuid4().hex[:6]}"

    dashboard = get_dashboard(db_session, company.name, user_roles="analyst", user_id=viewer)

    assert dashboard.periods[0].restricted
    assert {p.status for r in dashboard.ratios for p in r.points} == {"restricted"}
    assert all(p.value is None and not p.inputs for r in dashboard.ratios for p in r.points)
    session = SessionLocal()
    try:
        events = session.query(RetrievalAuditLog).filter_by(user_id=viewer).all()
    finally:
        session.close()
    assert events and all(not e.authorized and e.action == "view_dashboard" for e in events)


def test_unknown_entity_and_bad_period(db_session, company):
    company("EERR", 2024, INCOME[2024])
    with pytest.raises(LookupError):
        get_dashboard(db_session, f"Nobody{uuid.uuid4().hex[:6]}")
    with pytest.raises(DashboardError):
        refresh(db_session, company.name, periods=["sometime"], extractor=company.extractor)


def test_model_failure_is_reported(db_session, company):
    company("EERR", 2024, INCOME[2024])
    company.extractor.fail = True
    with pytest.raises(ExtractionServiceError):
        refresh(db_session, company.name, extractor=company.extractor)


# --- API -------------------------------------------------------------------------------------------


@pytest.fixture
def client():
    from app.main import app

    return TestClient(app)


def test_catalog_lists_entities_periods_and_statements_by_access(client, company):
    company("EERR", 2024, INCOME[2024])
    company("Balance", 2025, BALANCE[2025], tag="restricted:compensation")

    analyst = client.get("/dashboard/catalog", headers={"X-User-Role": "analyst"}).json()
    admin = client.get("/dashboard/catalog", headers={"X-User-Role": "admin"}).json()

    mine = next(e for e in analyst["entities"] if e["entity"] == company.name)
    assert mine["periods"] == ["2024"]
    assert [(d["statements"], d["periods"]) for d in mine["documents"]] == [(["income_statement"], ["2024"])]
    full = next(e for e in admin["entities"] if e["entity"] == company.name)
    assert full["periods"] == ["2024", "2025"]


def test_dashboard_endpoint_and_excel_export(client, db_session, company):
    company("EERR", 2024, INCOME[2024])
    company("Balance", 2024, BALANCE[2024])
    refresh(db_session, company.name, extractor=company.extractor)

    response = client.get("/dashboard", params={"entity": company.name, "categories": ["profitability"]})
    assert response.status_code == 200
    dashboard = response.json()
    assert {r["category"] for r in dashboard["ratios"]} == {"profitability"}

    exported = client.post("/dashboard/export.xlsx", json=dashboard)
    assert exported.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(exported.content))
    assert wb.sheetnames == ["Ratios", "Inputs and sources", "Outliers and notes"]
    rows = list(wb["Ratios"].iter_rows(values_only=True))
    assert rows[2] == ("Category", "Ratio", "Formula", "2024")
    roe = next(row for row in rows if row[1] == "Return on equity (ROE)")
    assert roe[3] == pytest.approx(650_000 / 5_000_000)
    sources = list(wb["Inputs and sources"].iter_rows(values_only=True))
    assert any(row[3] == "Utilidad neta" and row[6] == f"EERR_{company.name}_2024.xlsx" for row in sources)

    assert client.get("/dashboard", params={"entity": "Nobody-xyz"}).status_code == 404
    assert client.get("/dashboard", params={"entity": company.name, "categories": ["vibes"]}).status_code == 400


def test_periods_come_from_statement_line_items_not_working_columns():
    from types import SimpleNamespace

    from app.dashboard.catalog import describe_document

    def row(sheet, label, columns):
        values = "; ".join(f"{c}={v}" for c, v in columns)
        return SimpleNamespace(content_type="text", sheet_name=sheet, cell_range="A5:D5",
                               text=f"Sheet '{sheet}', {label}: A={label}; {values}")

    document = SimpleNamespace(id=uuid.uuid4(), filename="EEFF Eclipse 2024.xlsx", sheet_names=["ESF", "FC"],
                               confidentiality_tag="internal")
    chunks = [
        row("ESF", "Total activos", [("2024", 100), ("2023", 90)]),
        row("FC", "Metodo de participacion", [("2024", 5), ("2016", 7)]),  # a helper column
    ]
    described = describe_document(document, chunks)
    assert [p.label for p in described.periods] == ["2023", "2024"]
    assert described.statements == ["balance_sheet"]

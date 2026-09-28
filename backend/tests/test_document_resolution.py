"""Integration tests for ID-HU-FE-002's backend support: cell/formula lookup,
version supersession, and confidentiality-based access.

Requires a reachable database (run via `docker compose up`), same as
test_retrieval.py — skips otherwise. Route functions are called directly
(they're plain callables) rather than through a TestClient, matching this
codebase's existing service-level integration test style.
"""
import openpyxl
import pytest
from fastapi import HTTPException

from app.api.routes_documents import (
    delete_document,
    get_cell,
    get_document_file,
    list_documents,
    resolve_document,
    set_confidentiality,
    supersede_document,
)
from app.db.models import Document
from app.ingestion.service import ingest_excel_file
from app.schemas import ConfidentialityRequest, SupersedeRequest
from app.storage import save_document_file

# fake_embeddings and db_session fixtures come from conftest.py


def _make_workbook(tmp_path, name: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FCF"
    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    ws["A2"] = "Free Cash Flow"
    ws["B2"] = 42
    ws["C2"] = "=B2*1"
    file_path = tmp_path / name
    wb.save(file_path)
    return file_path


def test_get_cell_returns_value_and_formula(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model.xlsx")
    document = ingest_excel_file(db_session, filename="model.xlsx", file_path=file_path)
    try:
        # Ingested documents default to "restricted" (ID-HU-BE); this test is
        # about cell/formula resolution, not confidentiality, so classify it
        # as accessible to a plain analyst first.
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="public"), db=db_session)

        value_cell = get_cell(str(document.id), "FCF", "B2", db=db_session)
        assert value_cell.value == "42"
        assert value_cell.formula is None

        formula_cell = get_cell(str(document.id), "FCF", "C2", db=db_session)
        assert formula_cell.formula == "=B2*1"
        assert formula_cell.sheet_name == "FCF"
    finally:
        db_session.delete(document)
        db_session.commit()


def test_get_cell_unknown_sheet_returns_404(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model2.xlsx")
    document = ingest_excel_file(db_session, filename="model2.xlsx", file_path=file_path)
    try:
        # Accessible by default so the 404 below comes from the missing sheet,
        # not from the default-restricted confidentiality tag (ID-HU-BE).
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="public"), db=db_session)

        with pytest.raises(HTTPException) as exc_info:
            get_cell(str(document.id), "DoesNotExist", "A1", db=db_session)
        assert exc_info.value.status_code == 404
    finally:
        db_session.delete(document)
        db_session.commit()


def test_supersede_and_resolve_reports_current_version(tmp_path, db_session):
    old_path = _make_workbook(tmp_path, "budget_v3.xlsx")
    new_path = _make_workbook(tmp_path, "budget_v4.xlsx")
    old_doc = ingest_excel_file(db_session, filename="budget_v3.xlsx", file_path=old_path)
    new_doc = ingest_excel_file(db_session, filename="budget_v4.xlsx", file_path=new_path)
    try:
        supersede_document(str(old_doc.id), SupersedeRequest(new_document_id=new_doc.id), db=db_session)

        resolution = resolve_document(str(old_doc.id), db=db_session)
        assert resolution.is_superseded is True
        assert resolution.current_version is not None
        assert resolution.current_version.id == new_doc.id

        still_current = resolve_document(str(new_doc.id), db=db_session)
        assert still_current.is_superseded is False
        assert still_current.current_version is None
    finally:
        db_session.delete(old_doc)
        db_session.delete(new_doc)
        db_session.commit()


def test_get_document_file_serves_the_original_bytes(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model3.xlsx")
    document = ingest_excel_file(db_session, filename="model3.xlsx", file_path=file_path)
    try:
        # Mirrors what routes_ingest.py does after a successful ingestion.
        document.storage_path = save_document_file(document.id, ".xlsx", file_path)
        db_session.commit()
        # Accessible by default (ID-HU-BE default-restricts new documents);
        # this test only cares about file serving, not confidentiality.
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="public"), db=db_session)

        response = get_document_file(str(document.id), db=db_session)
        assert response.path.exists()
        assert response.path.read_bytes() == file_path.read_bytes()
    finally:
        db_session.delete(document)
        db_session.commit()


def test_get_document_file_404s_when_never_persisted(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model4.xlsx")
    document = ingest_excel_file(db_session, filename="model4.xlsx", file_path=file_path)
    try:
        # Accessible by default so the 404 below comes from the missing
        # storage path, not from the default-restricted confidentiality tag.
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="public"), db=db_session)

        with pytest.raises(HTTPException) as exc_info:
            get_document_file(str(document.id), db=db_session)
        assert exc_info.value.status_code == 404
    finally:
        db_session.delete(document)
        db_session.commit()


def test_get_document_file_denies_restricted_document(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model5.xlsx")
    document = ingest_excel_file(db_session, filename="model5.xlsx", file_path=file_path)
    try:
        document.storage_path = save_document_file(document.id, ".xlsx", file_path)
        db_session.commit()
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="restricted"), db=db_session)

        with pytest.raises(HTTPException) as exc_info:
            get_document_file(str(document.id), x_user_role="analyst", db=db_session)
        assert exc_info.value.status_code == 403
    finally:
        db_session.delete(document)
        db_session.commit()


def test_restricted_document_denies_analyst_but_allows_admin(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "comp_report.xlsx")
    document = ingest_excel_file(db_session, filename="comp_report.xlsx", file_path=file_path)
    try:
        set_confidentiality(str(document.id), ConfidentialityRequest(tag="restricted"), db=db_session)

        as_analyst = resolve_document(str(document.id), x_user_role="analyst", db=db_session)
        assert as_analyst.access.allowed is False
        assert as_analyst.access.reason

        as_admin = resolve_document(str(document.id), x_user_role="admin", db=db_session)
        assert as_admin.access.allowed is True

        with pytest.raises(HTTPException) as exc_info:
            get_cell(str(document.id), "FCF", "B2", x_user_role="analyst", db=db_session)
        assert exc_info.value.status_code == 403
    finally:
        db_session.delete(document)
        db_session.commit()


def test_delete_document_removes_it_and_nulls_out_dangling_supersession(tmp_path, db_session):
    old_path = _make_workbook(tmp_path, "budget_v5.xlsx")
    new_path = _make_workbook(tmp_path, "budget_v6.xlsx")
    old_doc = ingest_excel_file(db_session, filename="budget_v5.xlsx", file_path=old_path)
    new_doc = ingest_excel_file(db_session, filename="budget_v6.xlsx", file_path=new_path)
    supersede_document(str(old_doc.id), SupersedeRequest(new_document_id=new_doc.id), db=db_session)
    old_id = old_doc.id
    new_id = new_doc.id

    try:
        delete_document(str(new_id), db=db_session)

        assert db_session.get(Document, new_id) is None
        db_session.refresh(old_doc)
        assert old_doc.superseded_by_id is None

        with pytest.raises(HTTPException) as exc_info:
            resolve_document(str(new_id), db=db_session)
        assert exc_info.value.status_code == 404
    finally:
        db_session.delete(old_doc)
        db_session.commit()


def test_delete_document_unknown_id_404s(db_session):
    with pytest.raises(HTTPException) as exc_info:
        delete_document("00000000-0000-0000-0000-000000000000", db=db_session)
    assert exc_info.value.status_code == 404


def test_list_documents_includes_sheet_names(tmp_path, db_session):
    file_path = _make_workbook(tmp_path, "model6.xlsx")
    document = ingest_excel_file(db_session, filename="model6.xlsx", file_path=file_path)
    try:
        result = list_documents(db=db_session)
        listed = next(doc for doc in result.documents if doc.id == document.id)
        assert listed.sheet_names == ["FCF"]
    finally:
        db_session.delete(document)
        db_session.commit()

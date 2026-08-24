import openpyxl
import pytest

from app.ingestion.excel_parser import ExcelParsingError, parse_excel


def _cell(sheet, address):
    return next(c for c in sheet.cells if c.address == address)


# --- BE-001: sheet & cell value extraction ---


def test_extracts_sheet_and_cell_values(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Line Item"
    ws["B1"] = "Q1"
    ws["A2"] = "Revenue"
    ws["B2"] = 120000
    path = tmp_path / "q3_budget.xlsx"
    wb.save(path)

    result = parse_excel(path)

    assert len(result.sheets) == 1
    sheet = result.sheets[0]
    assert sheet.name == "Budget"
    assert sheet.is_empty is False
    assert _cell(sheet, "A2").value == "Revenue"
    assert _cell(sheet, "B2").value == "120000"


def test_merged_cell_value_associated_with_every_spanned_cell(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["B2"] = "Q3 Actuals"
    ws.merge_cells("B2:D2")
    path = tmp_path / "merged.xlsx"
    wb.save(path)

    result = parse_excel(path)
    sheet = result.sheets[0]

    for address in ("B2", "C2", "D2"):
        cell = _cell(sheet, address)
        assert cell.value == "Q3 Actuals"
        assert cell.is_merged is True
        assert cell.merged_range == "B2:D2"


def test_empty_sheet_indexed_without_error_and_rest_processed_normally(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Budget"
    ws1["A1"] = "Revenue"
    wb.create_sheet("Template")
    path = tmp_path / "with_empty_sheet.xlsx"
    wb.save(path)

    result = parse_excel(path)

    by_name = {s.name: s for s in result.sheets}
    assert by_name["Template"].is_empty is True
    assert by_name["Template"].cells == []
    assert by_name["Budget"].is_empty is False


# --- BE-002: formulas, named ranges, cross-sheet references, circular refs ---


def test_formula_and_cross_sheet_reference_preserved(tmp_path):
    wb = openpyxl.Workbook()
    assumptions = wb.active
    assumptions.title = "Assumptions"
    assumptions["B5"] = 0.1
    revenue = wb.create_sheet("Revenue")
    revenue["B2"] = 100
    revenue["B3"] = "=B2*(1+Assumptions!B5)"
    path = tmp_path / "projection.xlsx"
    wb.save(path)

    result = parse_excel(path)
    revenue_sheet = next(s for s in result.sheets if s.name == "Revenue")
    cell = _cell(revenue_sheet, "B3")

    assert cell.formula == "=B2*(1+Assumptions!B5)"
    assert "Assumptions!B5" in cell.formula_references


def test_named_range_resolved(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws["B5"] = 0.21
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Assumptions!$B$5")
    path = tmp_path / "named_range.xlsx"
    wb.save(path)

    result = parse_excel(path)

    named = next(n for n in result.named_ranges if n.name == "TaxRate")
    assert named.sheet_name == "Assumptions"
    assert named.cell_range == "$B$5"


def test_circular_reference_detected_and_ingestion_still_succeeds(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CashFlow"
    ws["A1"] = "=B1+1"
    ws["B1"] = "=A1+1"
    path = tmp_path / "circular.xlsx"
    wb.save(path)

    result = parse_excel(path)

    assert any("Circular reference" in w for w in result.warnings)
    assert len(result.sheets[0].cells) == 2


# --- BE-003: unsupported elements & warnings ---


def test_macro_workbook_flagged_without_execution(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "value"
    path = tmp_path / "with_macros.xlsm"
    wb.save(path)

    result = parse_excel(path)

    assert any("macro" in w.lower() for w in result.warnings)


def test_external_workbook_reference_flagged(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "=[Q2_Actuals.xlsx]Sheet1!A1"
    path = tmp_path / "external_link.xlsx"
    wb.save(path)

    result = parse_excel(path)

    assert any("Q2_Actuals.xlsx" in w for w in result.warnings)


def test_protected_sheet_flagged_not_fully_parsed(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Confidential"
    ws["A1"] = "secret"
    ws.protection.sheet = True
    path = tmp_path / "protected.xlsx"
    wb.save(path)

    result = parse_excel(path)
    sheet = result.sheets[0]

    assert sheet.is_protected is True
    assert any("protected" in w.lower() for w in result.warnings)


def test_corrupted_file_raises_clear_error(tmp_path):
    path = tmp_path / "corrupted.xlsx"
    path.write_bytes(b"not a real xlsx file")

    with pytest.raises(ExcelParsingError):
        parse_excel(path)

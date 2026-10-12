"""ID-HU-FE-003 — figure extraction and cross-document comparison.

Extraction tests ingest real workbooks (fake embeddings, see conftest.py)
and script the chat model; the comparison arithmetic is tested directly on
`summarize`, which involves no model at all.
"""
import re
import uuid
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException

import app.figures.extraction as extraction_module
from app.api.routes_compare import compare
from app.figures.comparison import compare_documents, summarize
from app.figures.export import comparison_workbook
from app.figures.extraction import (
    ExchangeRate,
    ExtractedFigure,
    ExtractionServiceError,
    LLMExchangeRate,
    LLMFigure,
    expected_coverage,
    extract_figure,
    orient_rate,
    single_currency,
    single_year,
)
from app.ingestion.service import ingest_excel_file
from app.qa.confidence import primary_reading
from app.qa.schemas import ResolvedCitation
from app.schemas import CompareRequest, ComparisonOut

# fake_embeddings and db_session fixtures come from conftest.py


class _FakeExtractionModel:
    """Plays the extraction model: for the document named in the prompt it
    returns the scripted figure, pointing at whichever source actually holds
    that number (as the real model would), or at `force_source` if given."""

    def __init__(self, figures: dict[str, dict], calls: list | None = None):
        self._figures = figures
        self.calls = calls if calls is not None else []

    def with_structured_output(self, schema):
        figures, calls = self._figures, self.calls

        class _Structured:
            def invoke(self, messages):
                prompt = messages[-1]["content"]
                calls.append(prompt)
                if schema is LLMExchangeRate:
                    return figures.get("__rate__", LLMExchangeRate(found=False))
                for filename, spec in figures.items():
                    if filename not in prompt:
                        continue
                    spec = dict(spec)
                    source_id = spec.pop("force_source", None)
                    if source_id is None:
                        digits = re.sub(r"\D", "", spec.get("value_text", ""))
                        for match in re.finditer(r"\[(S\d+)\][^\n]*\n([^\[]*)", prompt):
                            if digits and digits in re.sub(r"\D", "", match.group(2)):
                                source_id = match.group(1)
                                break
                    return LLMFigure(found=True, source_id=source_id, **spec)
                return LLMFigure(found=False)

        return _Structured()


def _patch(monkeypatch, figures, calls=None):
    model = _FakeExtractionModel(figures, calls)
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: model)
    return model


def _ingest(db_session, tmp_path, filename, rows, tag="public"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estado de resultados"
    for row in rows:
        ws.append(row)
    path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
    wb.save(path)
    document = ingest_excel_file(db_session, filename=filename, file_path=path)
    document.confidentiality_tag = tag
    for chunk in document.chunks:
        chunk.confidentiality_tag = tag
    db_session.commit()
    return document


@pytest.fixture
def cleanup(db_session):
    created = []
    yield created
    for document in created:
        db_session.delete(document)
    db_session.commit()


# --- number readings ----------------------------------------------------------

@pytest.mark.parametrize(
    ("token", "value"),
    [("4,820,000", 4820000), ("4.820.000", 4820000), ("81,000", 81000), ("1,234.5", 1234.5),
     ("1.234,5", 1234.5), ("12,5", 12.5), ("0.125", 0.125), ("4050.5", 4050.5), ("-7,52", -7.52)],
)
def test_primary_reading(token, value):
    assert primary_reading(token) == pytest.approx(value)


# --- extraction -------------------------------------------------------------------

def test_extracts_a_verified_figure_with_its_citation(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "EERR_Andina_2024_fe003.xlsx",
                  [["Concepto", "Valor"], ["Ingresos", 4820000], ["Utilidad neta", 610000]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {"value_text": "610,000", "currency": "cop", "period": "2024", "label": "Utilidad neta"}})

    figure = extract_figure(db_session, doc.id, "net income")

    assert figure.found is True
    assert figure.value == 610000
    assert figure.currency == "COP"
    assert figure.citation.document_id == str(doc.id)
    assert "610000" in figure.citation.text
    assert figure.confidence == "high"


def test_a_figure_not_in_its_source_is_not_found_rather_than_invented(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "EERR_fab_fe003.xlsx", [["Utilidad neta", 610000]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {"value_text": "999,999", "force_source": "S1"}})

    figure = extract_figure(db_session, doc.id, "net income")

    assert figure.found is False
    assert figure.value is None
    assert "couldn't be found in its source" in figure.reason


def test_restricted_document_is_not_extracted_for_an_analyst(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "EERR_restr_fe003.xlsx", [["Utilidad neta", 610000]], tag="restricted")
    cleanup.append(doc)
    model = _patch(monkeypatch, {doc.filename: {"value_text": "610,000"}})

    figure = extract_figure(db_session, doc.id, "net income", user_roles="analyst")

    assert figure.found is False
    assert "restricted" in figure.reason
    assert model.calls == []  # the model never saw the restricted content


def test_non_standard_label_and_scale(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "Balance_fe003.xlsx",
                  [["Cifras en miles"], ["Short-Term Obligations", 1250]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {
        "value_text": "1250", "unit_scale": "thousands", "label": "Short-Term Obligations",
        "label_matches_metric": False,
    }})

    figure = extract_figure(db_session, doc.id, "current liabilities")

    assert figure.value == 1250000
    assert figure.confidence == "medium"
    assert "Short-Term Obligations" in figure.reason


def test_a_scale_the_document_does_not_state_is_not_applied(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "BI_noscale_fe003.xlsx", [["Revenue", 4820000]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {"value_text": "4,820,000", "unit_scale": "thousands"}})

    figure = extract_figure(db_session, doc.id, "revenue")

    assert figure.value == 4820000
    assert figure.unit_scale == "units"
    assert figure.confidence == "medium"
    assert "doesn't say so" in figure.reason


# --- comparison arithmetic ----------------------------------------------------------

def _doc(name, current=True):
    return SimpleNamespace(id=uuid.uuid4(), filename=name, version_number=1, is_current=current)


def _fig(value, currency="COP", found=True):
    citation = ResolvedCitation("S1", "d", "f.xlsx", "Hoja", "A2:B2", None, None, f"Revenue: {value}")
    return ExtractedFigure(document_id="d", found=found, value=value if found else None,
                           value_text=f"{value:,}" if found else None, currency=currency,
                           citation=citation if found else None, confidence="high",
                           reason=None if found else "This document doesn't state revenue.")


def test_matching_values_reconcile():
    result = summarize("net income", None, [_doc("IS.xlsx"), _doc("CF.xlsx")],
                       [_fig(610000), _fig(610000)], {}, tolerance=0.001)
    assert result.reconciles is True
    assert result.rows[1].variance_abs == 0
    assert result.escalation is None


def test_discrepancy_is_flagged_with_variance_and_escalation():
    result = summarize("revenue", "Q1 2025", [_doc("BI_export.xlsx"), _doc("ERP_report.xlsx")],
                       [_fig(4820000), _fig(4795000)], {}, tolerance=0.001)

    assert result.reconciles is False
    erp = result.rows[1]
    assert erp.matches_baseline is False
    assert erp.variance_abs == -25000
    assert erp.variance_pct == pytest.approx(-0.5187, abs=1e-4)
    assert result.escalation.team == "Accounting/Consolidation"


def test_rounding_differences_within_tolerance_still_reconcile():
    result = summarize("revenue", None, [_doc("a.xlsx"), _doc("b.xlsx")],
                       [_fig(4820000), _fig(4820400)], {}, tolerance=0.001)
    assert result.reconciles is True


def _fig_in(value, period):
    figure = _fig(value)
    figure.period = period
    return figure


def test_variance_across_periods_is_shown_without_a_reconciliation_verdict():
    result = summarize("revenue", None, [_doc("2024.xlsx"), _doc("2025.xlsx"), _doc("2026.xlsx")],
                       [_fig_in(1000000, "FY2024"), _fig_in(1100000, "FY2025"), _fig_in(900000, "FY2026")],
                       {}, tolerance=0.001)

    assert [r.variance_pct for r in result.rows] == [None, pytest.approx(10.0), pytest.approx(-10.0)]
    assert result.rows[1].variance_abs == 100000
    assert result.rows[0].is_baseline is True
    assert result.across_periods is True
    assert result.reconciles is None
    assert all(r.matches_baseline is None for r in result.rows)
    assert result.escalation is None


def test_same_period_written_differently_is_still_a_reconciliation():
    result = summarize("revenue", None, [_doc("a.xlsx"), _doc("b.xlsx")],
                       [_fig_in(4820000, "Q1 2025"), _fig_in(4795000, "q1  2025")], {}, tolerance=0.001)
    assert result.across_periods is False
    assert result.reconciles is False


def test_a_missing_figure_is_noted_and_cannot_reconcile_alone():
    result = summarize("revenue", None, [_doc("a.xlsx"), _doc("b.xlsx")],
                       [_fig(100000), _fig(0, found=False)], {}, tolerance=0.001)
    assert result.reconciles is None
    assert any("b.xlsx" in note for note in result.notes)


def test_other_currency_is_converted_with_a_documented_rate():
    citation = ResolvedCitation("S1", "d", "fx.xlsx", None, None, None, None, "TRM USD/COP 4,000")
    rate = ExchangeRate("USD", "COP", 4000.0, "4,000", citation)
    result = summarize("revenue", None, [_doc("co.xlsx"), _doc("us.xlsx")],
                       [_fig(4000000, "COP"), _fig(1000, "USD")], {"USD": rate}, tolerance=0.001)

    us = result.rows[1]
    assert us.converted is True
    assert us.comparable_value == 4000000
    assert us.value == 1000  # original currency value kept
    assert result.reconciles is True
    assert result.exchange_rates == [rate]


def test_other_currency_without_documented_rate_is_not_converted():
    result = summarize("revenue", None, [_doc("eu.xlsx"), _doc("us.xlsx")],
                       [_fig(1000, "EUR"), _fig(1100, "USD")], {}, tolerance=0.001)
    assert result.reconciles is None
    assert result.rows[1].comparable_value is None
    assert any("No documented exchange rate" in note for note in result.notes)


def test_superseded_versions_are_noted():
    result = summarize("revenue", None, [_doc("v1.xlsx", current=False), _doc("v2.xlsx")],
                       [_fig(1), _fig(1)], {}, tolerance=0.001)
    assert any("not the current version" in note for note in result.notes)


@pytest.mark.parametrize(
    ("base", "quote", "expected"),
    [("USD", "COP", 4000.0), ("COP", "USD", 1 / 4000.0), ("EUR", "GBP", None)],
)
def test_orient_rate(base, quote, expected):
    rate = 4000.0 if base == "USD" else (1 / 4000.0 if base == "COP" else 0.85)
    oriented = orient_rate(rate, base, quote, "USD", "COP")
    assert oriented == (pytest.approx(4000.0) if expected else None)


# --- end to end (fake model) and API ------------------------------------------------

def test_compare_documents_end_to_end(monkeypatch, db_session, tmp_path, cleanup):
    bi = _ingest(db_session, tmp_path, "BI_export_fe003.xlsx", [["Revenue Q1 2025", 4820000]])
    erp = _ingest(db_session, tmp_path, "ERP_report_fe003.xlsx", [["Revenue Q1 2025", 4795000]])
    cleanup += [bi, erp]
    _patch(monkeypatch, {
        bi.filename: {"value_text": "4,820,000", "currency": "COP", "period": "Q1 2025"},
        erp.filename: {"value_text": "4,795,000", "currency": "COP", "period": "Q1 2025"},
    })

    result = compare_documents(db_session, [bi, erp], "revenue")

    assert [r.value for r in result.rows] == [4820000, 4795000]
    assert result.reconciles is False
    assert all(r.citation is not None for r in result.rows)


def test_compare_rejects_selecting_a_document_twice(db_session):
    document_id = uuid.uuid4()
    with pytest.raises(HTTPException) as exc_info:
        compare(CompareRequest(document_ids=[document_id, document_id], metric="revenue"), db=db_session)
    assert exc_info.value.status_code == 400


def test_compare_request_needs_at_least_two_documents():
    with pytest.raises(ValueError):
        CompareRequest(document_ids=[uuid.uuid4()], metric="revenue")


def test_excel_export_has_values_citations_and_the_discrepancy():
    result = summarize("revenue", "Q1 2025", [_doc("BI_export.xlsx"), _doc("ERP_report.xlsx")],
                       [_fig(4820000), _fig(4795000)], {}, tolerance=0.001)
    workbook = openpyxl.load_workbook(BytesIO(comparison_workbook(ComparisonOut.model_validate(result))))
    ws = workbook.active
    values = [[cell.value for cell in row] for row in ws.iter_rows()]

    flat = [str(v) for row in values for v in row if v is not None]
    assert "Does NOT reconcile" in flat
    assert "BI_export.xlsx" in flat and "ERP_report.xlsx" in flat
    assert any("sheet 'Hoja'" in v for v in flat)              # citation column
    assert any(row[11] == "NO" for row in values if len(row) > 11)
    assert "Accounting/Consolidation" in flat


# --- period coverage (three months vs year-to-date) and model failures ---------------

@pytest.mark.parametrize(
    ("period", "expected"),
    [("Q3 2026", {"three_months", "point_in_time"}), ("3T 2025", {"three_months", "point_in_time"}),
     ("FY2024", {"twelve_months", "point_in_time"}), ("2024", {"twelve_months", "point_in_time"}),
     ("H1 2025", {"six_months", "point_in_time"}), (None, None), ("June 2026", None)],
)
def test_expected_coverage(period, expected):
    assert expected_coverage(period) == expected


def test_year_to_date_value_for_a_quarter_is_flagged(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "10Q_coverage_fe003.xlsx",
                  [["Net income", "Three Months", "Nine Months"], ["Net income", 29789, 101464]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {
        "value_text": "101,464", "period": "Q3 2026", "label": "Net income",
        "period_coverage": "nine_months", "column_header": "Nine Months",
    }})

    figure = extract_figure(db_session, doc.id, "net income", period="Q3 2026")

    assert figure.found and figure.confidence == "medium"
    assert figure.period_coverage == "nine_months"
    assert "nine months" in figure.issues["coverage"]
    assert "Nine Months" in figure.reason


def test_matching_coverage_keeps_high_confidence(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "10Q_coverage_ok_fe003.xlsx",
                  [["Net income", "Three Months", "Nine Months"], ["Net income", 29789, 101464]])
    cleanup.append(doc)
    _patch(monkeypatch, {doc.filename: {
        "value_text": "29,789", "period": "Q3 2026", "label": "Net income", "period_coverage": "three_months",
    }})

    figure = extract_figure(db_session, doc.id, "net income", period="Q3 2026")

    assert figure.confidence == "high" and not figure.issues


class _BrokenModel:
    def with_structured_output(self, schema):
        class _Structured:
            def invoke(self, messages):
                raise ConnectionError("model service down")
        return _Structured()


def test_model_failure_is_reported_as_an_error_not_as_missing(monkeypatch, db_session, tmp_path, cleanup):
    doc = _ingest(db_session, tmp_path, "BI_down_fe003.xlsx", [["Revenue", 4820000]])
    cleanup.append(doc)
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: _BrokenModel())

    figure = extract_figure(db_session, doc.id, "revenue")

    assert not figure.found and figure.error
    assert "didn't respond" in figure.reason


def test_compare_answers_503_when_the_model_is_down(monkeypatch, db_session, tmp_path, cleanup):
    a = _ingest(db_session, tmp_path, "BI_down_a_fe003.xlsx", [["Revenue", 4820000]])
    b = _ingest(db_session, tmp_path, "BI_down_b_fe003.xlsx", [["Revenue", 4795000]])
    cleanup += [a, b]
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: _BrokenModel())

    with pytest.raises(ExtractionServiceError):
        compare_documents(db_session, [a, b], "revenue")
    with pytest.raises(HTTPException) as exc_info:
        compare(CompareRequest(document_ids=[a.id, b.id], metric="revenue"), db=db_session)
    assert exc_info.value.status_code == 503


def test_same_label_different_coverage_is_compared_as_different_periods():
    quarter, year_to_date = _fig(29789), _fig(101464)
    quarter.period, year_to_date.period = "Q3 2026", "Q3 2026"
    quarter.period_coverage, year_to_date.period_coverage = "three_months", "nine_months"

    result = summarize("net income", "Q3 2026", [_doc("IS.xlsx"), _doc("10Q.pdf")],
                       [quarter, year_to_date], {}, tolerance=0.001)

    assert result.across_periods is True
    assert result.reconciles is None and result.escalation is None
    assert any("different lengths of time" in note for note in result.notes)


def test_difference_with_an_unknown_period_is_not_called_a_discrepancy():
    # 2024 vs 2025 net income where extraction couldn't tell the periods apart.
    result = summarize("net income", None, [_doc("EERR_2024.xlsx"), _doc("EERR_2025.xlsx")],
                       [_fig(610000), _fig(702000)], {}, tolerance=0.001)

    assert result.reconciles is None
    assert result.escalation is None
    assert result.rows[1].matches_baseline is None
    assert result.rows[1].variance_abs == 92000  # the variance is still shown
    assert any("couldn't be determined" in note for note in result.notes)


def test_a_requested_period_still_flags_a_discrepancy_without_extracted_periods():
    result = summarize("revenue", "2024", [_doc("EERR_2024.xlsx"), _doc("BI_2024.xlsx")],
                       [_fig(4820000), _fig(4795000)], {}, tolerance=0.001)

    assert result.reconciles is False
    assert result.escalation is not None


@pytest.mark.parametrize(
    ("text", "value_text", "expected"),
    [("Sheet 'EERR' (Año 2024), Utilidad neta: Concepto=Utilidad neta; 2024=610000", "610,000", "2024"),
     ("Revenue: 2024=4820000; 2025=5310000", "4,820,000", None),  # two years: ambiguous
     ("Revenue: B=4820000", "4,820,000", None),
     ("Sheet 'Sales' (FY2023), Units: 2010", "2010", "2023")],  # the value isn't read as a year
)
def test_single_year_fallback(text, value_text, expected):
    assert single_year(text, value_text) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [("Sheet 'Balance' (Cifras en pesos colombianos (COP)), Total activos: 2024=6100000", "COP"),
     ("Sheet 'Summary' (Cifras en dólares estadounidenses (USD)), Revenue: 2024=1205", "USD"),
     ("Revenue: B=4820000", None),
     ("Tasa de cambio USD/COP: 4000", None)],  # two currencies: ambiguous
)
def test_single_currency_fallback(text, expected):
    assert single_currency(text) == expected


@pytest.mark.parametrize(
    ("a", "b", "same"),
    [("As of June 30, 2026", "June 30, 2026", True), ("FY2024", "2024", True), ("FY 2024", "2024", True),
     ("Quarter Ended June 30 2026", "June 30, 2026", True), ("Q1 2025", "Q1 2025", True),
     ("2024", "2025", False), ("Q1 2025", "Q2 2025", False)],
)
def test_period_key(a, b, same):
    from app.figures.extraction import period_key

    assert (period_key(a) == period_key(b)) is same


@pytest.mark.parametrize(
    ("a", "b", "same"),
    [("June 30, 2026", "Q2 2026", True), ("As of June 30, 2026", "Q2 2026", True),
     ("30 de junio de 2026", "Q2 2026", True), ("December 31, 2025", "FY2025", True),
     ("31 de diciembre de 2024", "2024", True), ("September 30, 2026", "Q3 2026", True),
     ("June 30, 2026", "Q2 2025", False), ("June 15, 2026", "Q2 2026", False), ("Q2 2026", "2026", False)],
)
def test_quarter_end_dates_are_the_quarter(a, b, same):
    from app.figures.extraction import same_period

    assert same_period(a, b) is same


def test_distinct_periods_groups_names_of_one_period_and_skips_unknown():
    from app.figures.extraction import distinct_periods

    assert distinct_periods(["As of June 30, 2026", None, "Q2 2026", "June 30, 2026"]) == ["As of June 30, 2026"]
    assert distinct_periods(["2024", "FY2025"]) == ["2024", "FY2025"]


@pytest.mark.parametrize(
    ("value_text", "source", "column", "coverage"),
    [("364,357", "row 'Total net sales': Three Months Ended June 27, 2026=109,417; Nine Months Ended June 27, 2026=364,357",
      "Nine Months Ended June 27, 2026", "nine_months"),
     ("109,417", "row 'Total net sales': Three Months Ended June 27, 2026=109,417; Nine Months Ended June 27, 2026=364,357",
      "Three Months Ended June 27, 2026", "three_months"),
     ("39,544", "row 'Cash': June 27, 2026=39,544; September 27, 2025=35,934", "June 27, 2026", "point_in_time"),
     ("343,524", "row 'Total current assets': As of June 30, 2026=343,524", "As of June 30, 2026", "point_in_time"),
     ("610000", "Sheet 'EERR', Utilidad neta: Concepto=Utilidad neta; 2024=610000", "2024", None),
     ("100", "row 'X': column 1=100; column 2=100", None, None),
     ("166,835", "row 'Cost of sales': column 7=54,647; column 8=166,835", None, None)],
)
def test_column_and_coverage_come_from_the_source_row(value_text, source, column, coverage):
    from app.figures.extraction import column_of, coverage_of

    assert column_of(value_text, source) == (column if column != "column 1" else None)
    assert coverage_of(column_of(value_text, source)) == coverage


def test_source_row_label():
    from app.figures.extraction import source_row_label

    assert source_row_label("Financial table page 4 (…), row 'Cost of sales: Products': Three Months=47,153") == "Cost of sales: Products"
    assert source_row_label("Sheet 'EERR', Utilidad neta: Concepto=Utilidad neta") is None

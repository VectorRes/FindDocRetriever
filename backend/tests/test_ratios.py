"""ID-HU-BE-010 — financial ratio calculation.

Formula, mapping and engine tests run on scripted figures (no database, no
model). The end-to-end test ingests real workbooks (fake embeddings, see
conftest.py) and scripts the extraction model, so the whole path —
retrieval, figure verification, mapping, arithmetic, citations — runs for real.
"""
import re
import uuid
from dataclasses import dataclass, field

import openpyxl
import pytest
from fastapi import HTTPException

import app.figures.extraction as extraction_module
from app.api.routes_ratios import compute_ratios, list_ratio_definitions
from app.figures.extraction import ExtractedFigure, ExtractionServiceError, LLMFigure
from app.ingestion.service import ingest_excel_file
from app.qa.schemas import ResolvedCitation
from app.ratios.engine import (
    calculate_ratios,
    compute_ratio,
    map_line_items,
    order_documents,
    select_ratios,
    to_mapped_input,
)
from app.ratios.formulas import CATEGORIES, LINE_ITEMS, RATIOS
from app.ratios.mapping import label_matches, normalize_label, statement_hint_score
from app.schemas import RatioReportOut, RatioRequest

# fake_embeddings and db_session fixtures come from conftest.py

# Sample statements (COP, FY2024). Income statement and balance sheet are
# separate documents, as in the cross-document alternate case.
INCOME_STATEMENT = {
    "revenue": ("Revenue", 4_820_000),
    "cost_of_sales": ("Cost of sales", -2_892_000),
    "gross_profit": ("Gross profit", 1_928_000),
    "operating_income": ("Operating income", 964_000),
    "ebitda": ("EBITDA", 1_205_000),
    "interest_expense": ("Interest expense", -120_500),
    "net_income": ("Net income", 610_000),
}
BALANCE_SHEET = {
    "cash": ("Cash and cash equivalents", 450_000),
    "accounts_receivable": ("Accounts receivable", 803_000),
    "inventory": ("Inventory", 578_400),
    "current_assets": ("Total current assets", 2_100_000),
    "total_assets": ("Total assets", 6_100_000),
    "current_liabilities": ("Total current liabilities", 1_200_000),
    "total_liabilities": ("Total liabilities", 3_660_000),
    "total_equity": ("Total equity", 2_440_000),
}

EXPECTED = {
    "current_ratio": 2_100_000 / 1_200_000,
    "quick_ratio": (2_100_000 - 578_400) / 1_200_000,
    "cash_ratio": 450_000 / 1_200_000,
    "debt_to_equity": 3_660_000 / 2_440_000,
    "debt_ratio": 3_660_000 / 6_100_000,
    "interest_coverage": 964_000 / 120_500,
    "gross_margin": 1_928_000 / 4_820_000,
    "operating_margin": 964_000 / 4_820_000,
    "ebitda_margin": 1_205_000 / 4_820_000,
    "net_margin": 610_000 / 4_820_000,
    "roa": 610_000 / 6_100_000,
    "roe": 610_000 / 2_440_000,
    "asset_turnover": 4_820_000 / 6_100_000,
    "inventory_turnover": 2_892_000 / 578_400,
    "receivables_turnover": 4_820_000 / 803_000,
}


@dataclass(eq=False)  # hashable by identity, like ORM rows
class _Doc:
    filename: str
    sheet_names: list = field(default_factory=list)
    is_current: bool = True
    id: uuid.UUID = field(default_factory=uuid.uuid4)


def _document(filename, sheet_names=(), is_current=True):
    return _Doc(filename, list(sheet_names), is_current)


def _citation(document, label, value):
    return ResolvedCitation(
        source_id="S1", document_id=str(document.id), document_filename=document.filename,
        sheet_name="Sheet1", cell_range=f"A{abs(hash(label)) % 50 + 1}", page_number=None,
        reference_number=None, text=f"{label}: {value:,}",
    )


class FakeExtractor:
    """Plays extract_figure from per-document tables {item_key: (label, value)}."""

    def __init__(self, tables, currency="COP", period="FY2024", periods=None, currencies=None, overrides=None):
        self.tables = tables  # {document: {key: (label, value)}}
        self.currency, self.period = currency, period
        self.periods = periods or {}
        self.currencies = currencies or {}
        self.overrides = overrides or {}  # {key: {ExtractedFigure field: value}}
        self.calls = []

    def __call__(self, db, document_id, metric, user_roles=None, user_id=None, period=None, **_):
        self.calls.append((document_id, metric))
        document = next(d for d in self.tables if d.id == document_id)
        key = next(k for k, item in LINE_ITEMS.items() if item.query == metric)
        if key not in self.tables[document]:
            return ExtractedFigure(document_id=str(document_id), found=False, reason=f"This document doesn't state {metric}.")
        label, value = self.tables[document][key]
        figure = ExtractedFigure(
            document_id=str(document_id), found=True, value=float(value), value_text=f"{value:,}",
            currency=self.currencies.get(key, self.currency), period=self.periods.get(key, self.period),
            label=label, citation=_citation(document, label, value), confidence="high",
        )
        for field_name, field_value in self.overrides.get(key, {}).items():
            setattr(figure, field_name, field_value)
        return figure


@pytest.fixture
def statements():
    income = _document("EERR_Andina_2024.xlsx", ["Estado de resultados"])
    balance = _document("Balance_Andina_2024.xlsx", ["Balance general"])
    return income, balance


def _run(tables, documents, **kwargs):
    extractor = FakeExtractor(tables, **{k: kwargs.pop(k) for k in ("periods", "currencies", "overrides") if k in kwargs})
    report = calculate_ratios(None, documents, extractor=extractor, **kwargs)
    return report, {r.key: r for r in report.ratios}, extractor


# --- formula library ----------------------------------------------------------

def test_library_covers_the_four_categories_and_the_named_ratios():
    assert {r.category for r in RATIOS.values()} == set(CATEGORIES)
    assert {"roe", "roa", "ebitda_margin", "current_ratio", "debt_to_equity"} <= RATIOS.keys()
    for ratio in RATIOS.values():
        assert set(ratio.inputs) <= LINE_ITEMS.keys(), ratio.key


@pytest.mark.parametrize("key", sorted(EXPECTED))
def test_each_formula_against_the_sample_statements(key):
    values = {k: float(v) for k, (_, v) in {**INCOME_STATEMENT, **BALANCE_SHEET}.items()}
    assert RATIOS[key].compute(values) == pytest.approx(EXPECTED[key])


def test_definitions_endpoint_lists_the_library():
    definitions = list_ratio_definitions()
    assert [d.key for d in definitions] == list(RATIOS)
    assert next(d for d in definitions if d.key == "roa").inputs == ["net_income", "total_assets"]


def test_select_ratios_by_key_and_category():
    assert [r.key for r in select_ratios(["roa", "ebitda_margin"])] == ["ebitda_margin", "roa"]
    assert {r.category for r in select_ratios(categories=["liquidity"])} == {"liquidity"}
    with pytest.raises(ValueError):
        select_ratios(["made_up_ratio"])


# --- line-item mapping ----------------------------------------------------------

@pytest.mark.parametrize(
    ("key", "label"),
    [("current_liabilities", "Total current liabilities"), ("current_liabilities", "Pasivos corrientes"),
     ("current_liabilities", "TOTAL PASIVO CORRIENTE (COP)"), ("net_income", "Utilidad neta del ejercicio 2024"),
     ("revenue", "Ingresos operacionales"), ("total_equity", "Shareholders' Equity"), ("cash", "Efectivo")],
)
def test_standard_labels_map_with_confidence(key, label):
    assert label_matches(LINE_ITEMS[key], label)


@pytest.mark.parametrize(
    ("key", "label"),
    [("current_liabilities", "Short-Term Obligations"), ("current_liabilities", "Non-current liabilities"),
     ("total_assets", "Total current assets"), ("net_income", "Net income attributable to minority interest"),
     ("revenue", None)],
)
def test_non_standard_labels_are_not_accepted(key, label):
    assert not label_matches(LINE_ITEMS[key], label)


def test_normalize_label_drops_accents_years_and_total():
    assert normalize_label("Total Pasivos Corrientes (Nota 12) 2024") == "pasivo corriente"
    assert normalize_label("Situación financiera") == "situacion financiera"
    # Singular and plural compare equal.
    assert normalize_label("Efectivo y equivalente de efectivo") == normalize_label("Efectivo y equivalentes de efectivo")
    assert normalize_label("Inventories") == normalize_label("Inventory")


def test_statement_hints_order_documents(statements):
    income, balance = statements
    assert statement_hint_score("income_statement", [income.filename, *income.sheet_names]) > 0
    assert statement_hint_score("balance_sheet", [balance.filename]) > 0
    assert order_documents([income, balance], LINE_ITEMS["total_assets"])[0] is balance
    assert order_documents([income, balance], LINE_ITEMS["net_income"])[0] is income


def test_non_standard_label_lowers_mapping_confidence(statements):
    _, balance = statements
    figure = ExtractedFigure(
        document_id=str(balance.id), found=True, value=1_200_000.0, value_text="1,200,000",
        label="Short-Term Obligations", citation=_citation(balance, "Short-Term Obligations", 1_200_000),
        confidence="high",
    )
    mapped = to_mapped_input(LINE_ITEMS["current_liabilities"], balance, figure)
    assert mapped.needs_review and mapped.confidence == "low"
    assert "Short-Term Obligations" in mapped.reason


# --- main case: EBITDA margin with citations --------------------------------------

def test_ebitda_margin_is_calculated_with_citations_to_revenue_and_ebitda(statements):
    income, _ = statements
    report, ratios, _ = _run({income: INCOME_STATEMENT}, [income], ratio_keys=["ebitda_margin"])
    ebitda_margin = ratios["ebitda_margin"]
    assert ebitda_margin.status == "calculated"
    assert ebitda_margin.value == pytest.approx(0.25)
    assert ebitda_margin.unit == "percent"
    assert ebitda_margin.period == "FY2024" and ebitda_margin.currency == "COP"
    assert not ebitda_margin.cross_document
    cited = {c.text for c in ebitda_margin.citations}
    assert cited == {"EBITDA: 1,205,000", "Revenue: 4,820,000"}


def test_all_categories_are_calculated_from_income_statement_and_balance_sheet(statements):
    income, balance = statements
    report, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance])
    assert set(ratios) == set(RATIOS)
    for key, expected in EXPECTED.items():
        assert ratios[key].status == "calculated", (key, ratios[key].reasons)
        assert ratios[key].value == pytest.approx(expected), key
        # A citation for every line item used.
        assert len(ratios[key].citations) == len(RATIOS[key].inputs), key
    assert not report.notes


# --- alternate case 1: cross-document ratio ----------------------------------------

def test_roa_combines_income_statement_and_balance_sheet(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"])
    roa = ratios["roa"]
    assert roa.status == "calculated"
    assert roa.value == pytest.approx(0.1)
    assert roa.cross_document
    assert {c.document_filename for c in roa.citations} == {income.filename, balance.filename}
    by_key = {i.key: i for i in roa.inputs}
    assert by_key["net_income"].document_filename == income.filename
    assert by_key["total_assets"].document_filename == balance.filename


def test_each_input_is_looked_up_in_the_likeliest_document_first(statements):
    income, balance = statements
    _, _, extractor = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"])
    # One call per input: the first document tried already had it with high confidence.
    assert extractor.calls == [
        (income.id, LINE_ITEMS["net_income"].query),
        (balance.id, LINE_ITEMS["total_assets"].query),
    ]


def test_input_missing_from_every_document_makes_the_ratio_not_calculable(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: {}}, [income, balance], ratio_keys=["roa", "net_margin"])
    assert ratios["roa"].status == "not_calculable" and ratios["roa"].value is None
    assert "Total assets" in ratios["roa"].reasons[0]
    assert ratios["net_margin"].status == "calculated"


def test_be011_missing_current_liabilities_is_named_as_the_missing_input(statements):
    # ID-HU-BE-011 alternate 1: the balance sheet has no current liabilities.
    _, balance = statements
    table = {k: v for k, v in BALANCE_SHEET.items() if k != "current_liabilities"}
    _, ratios, _ = _run({balance: table}, [balance], ratio_keys=["current_ratio", "debt_ratio"])
    current_ratio = ratios["current_ratio"]
    assert current_ratio.status == "not_calculable" and current_ratio.value is None
    assert current_ratio.missing_inputs == ["Current liabilities"]
    assert current_ratio.reasons == ["Missing input: Current liabilities."]
    assert ratios["debt_ratio"].status == "calculated" and ratios["debt_ratio"].missing_inputs == []


def test_be011_not_calculable_for_another_reason_lists_no_missing_input(statements):
    _, balance = statements
    table = {**BALANCE_SHEET, "current_liabilities": ("Total current liabilities", 0)}
    _, ratios, _ = _run({balance: table}, [balance], ratio_keys=["current_ratio"])
    assert ratios["current_ratio"].status == "not_calculable"
    assert ratios["current_ratio"].missing_inputs == []


# --- alternate case 2: ambiguous line-item mapping ------------------------------------

def test_non_standard_label_holds_the_ratio_for_review(statements):
    _, balance = statements
    table = {**BALANCE_SHEET, "current_liabilities": ("Short-Term Obligations", 1_200_000)}
    report, ratios, _ = _run({balance: table}, [balance], ratio_keys=["current_ratio", "debt_ratio"])
    current_ratio = ratios["current_ratio"]
    assert current_ratio.status == "needs_review"
    assert current_ratio.value == pytest.approx(1.75)
    assert any("Short-Term Obligations" in reason for reason in current_ratio.reasons)
    flagged = next(i for i in current_ratio.inputs if i.key == "current_liabilities")
    assert flagged.needs_review and flagged.confidence == "low"
    # Ratios that don't use the flagged input aren't held back.
    assert ratios["debt_ratio"].status == "calculated"
    assert any("review" in note for note in report.notes)


def test_a_standard_label_elsewhere_beats_a_non_standard_one(statements):
    income, balance = statements
    tables = {
        income: {"current_liabilities": ("Short-Term Obligations", 999_999), "current_assets": ("Current assets", 2_100_000)},
        balance: BALANCE_SHEET,
    }
    _, ratios, _ = _run(tables, [income, balance], ratio_keys=["current_ratio"])
    assert ratios["current_ratio"].status == "calculated"
    assert ratios["current_ratio"].value == pytest.approx(1.75)


# --- consistency checks ----------------------------------------------------------------

def test_inputs_from_different_periods_need_review(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance],
                        ratio_keys=["roa"], periods={"total_assets": "FY2023"})
    assert ratios["roa"].status == "needs_review"
    assert "different periods" in ratios["roa"].reasons[0]


def test_inputs_in_different_currencies_are_not_divided(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance],
                        ratio_keys=["roa"], currencies={"total_assets": "USD"})
    assert ratios["roa"].status == "not_calculable" and ratios["roa"].value is None
    assert "different currencies" in ratios["roa"].reasons[0]


def test_zero_denominator_is_reported_not_raised(statements):
    _, balance = statements
    table = {**BALANCE_SHEET, "current_liabilities": ("Total current liabilities", 0)}
    _, ratios, _ = _run({balance: table}, [balance], ratio_keys=["current_ratio"])
    assert ratios["current_ratio"].status == "not_calculable"
    assert "denominator is zero" in ratios["current_ratio"].reasons[0]


def test_implausible_result_is_held_for_review(statements):
    # Nine-month net income over one quarter's revenue: a margin above 100%.
    income, _ = statements
    table = {**INCOME_STATEMENT, "net_income": ("Net income", 9_640_000)}
    _, ratios, _ = _run({income: table}, [income], ratio_keys=["net_margin"])
    assert ratios["net_margin"].status == "needs_review"
    assert ratios["net_margin"].value == pytest.approx(2.0)
    assert "plausible range" in ratios["net_margin"].reasons[0]


def test_scale_applied_to_only_some_figures_of_one_document_cancels_out(statements):
    # Extraction read "in millions" for one figure of the balance sheet but not the other.
    _, balance = statements
    _, ratios, _ = _run({balance: BALANCE_SHEET}, [balance], ratio_keys=["current_ratio"], overrides={
        "current_assets": {"value": 2_100_000 * 1e6, "unit_scale": "millions"},
    })
    assert ratios["current_ratio"].status == "calculated"
    assert ratios["current_ratio"].value == pytest.approx(1.75)


def test_documented_scales_are_kept_across_documents(statements):
    # Income statement in thousands, balance sheet in units: base values are compared.
    income, balance = statements
    tables = {income: {"net_income": ("Net income", 610)}, balance: BALANCE_SHEET}
    _, ratios, _ = _run(tables, [income, balance], ratio_keys=["roa"], overrides={
        "net_income": {"value": 610_000.0, "unit_scale": "thousands"},
    })
    assert ratios["roa"].value == pytest.approx(0.1)


def test_medium_confidence_input_is_shown_without_holding_the_ratio(statements):
    income, _ = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT}, [income], ratio_keys=["net_margin"], overrides={
        "net_income": {"confidence": "medium", "reason": "The document also has FY2023; the most recent period was used."},
    })
    assert ratios["net_margin"].status == "calculated"
    assert any("FY2023" in reason for reason in ratios["net_margin"].reasons)


def test_mixed_periods_without_a_requested_period_suggest_one(statements):
    income, balance = statements
    report, _, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance],
                        ratio_keys=["roa"], periods={"total_assets": "Q3 2026"})
    assert any("Specify a period" in note for note in report.notes)


def test_compute_ratio_is_pure_over_mapped_inputs(statements):
    income, balance = statements
    mapped = map_line_items(None, [income, balance], ["net_income", "total_equity"],
                            extractor=FakeExtractor({income: INCOME_STATEMENT, balance: BALANCE_SHEET}))
    assert compute_ratio(RATIOS["roe"], mapped).value == pytest.approx(0.25)


def test_superseded_document_is_noted(statements):
    income, _ = statements
    old = _document("EERR_Andina_2024_v1.xlsx", is_current=False)
    report, _, _ = _run({old: INCOME_STATEMENT}, [old], ratio_keys=["net_margin"])
    assert any("not the current version" in note for note in report.notes)


# --- API --------------------------------------------------------------------------------

def test_ratio_request_needs_a_document():
    with pytest.raises(ValueError):
        RatioRequest(document_ids=[])


def test_route_rejects_selecting_a_document_twice(db_session):
    document_id = uuid.uuid4()
    with pytest.raises(HTTPException) as exc:
        compute_ratios(RatioRequest(document_ids=[document_id, document_id]), db=db_session)
    assert exc.value.status_code == 400


def test_route_rejects_an_unknown_ratio(monkeypatch, db_session, tmp_path, cleanup):
    document = _ingest(db_session, tmp_path, "EERR_unknown_ratio_be010.xlsx", [["Concepto", "Valor"], ["Ingresos", 1000]])
    cleanup.append(document)
    with pytest.raises(HTTPException) as exc:
        compute_ratios(RatioRequest(document_ids=[document.id], ratios=["made_up_ratio"]), db=db_session)
    assert exc.value.status_code == 400


# --- end to end: real ingestion and figure verification, scripted model -------------------

class _FakeExtractionModel:
    """Answers the extraction prompt from {filename: {metric_query: (label, value_text)}},
    pointing at the source that actually holds the number."""

    def __init__(self, figures):
        self.figures = figures

    def with_structured_output(self, schema):
        figures = self.figures

        class _Structured:
            def invoke(self, messages):
                prompt = messages[-1]["content"]
                metric = re.search(r"METRIC: (.*)", prompt).group(1)
                for filename, items in figures.items():
                    if filename not in prompt or metric not in items:
                        continue
                    label, value_text = items[metric]
                    digits = re.sub(r"\D", "", value_text)
                    source_id = next(
                        (m.group(1) for m in re.finditer(r"\[(S\d+)\][^\n]*\n([^\[]*)", prompt)
                         if digits in re.sub(r"\D", "", m.group(2))),
                        None,
                    )
                    return LLMFigure(found=True, source_id=source_id, value_text=value_text,
                                     currency="COP", period="2024", label=label)
                return LLMFigure(found=False)

        return _Structured()


def _ingest(db_session, tmp_path, filename, rows, sheet="Hoja1", tag="public"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
    wb.save(path)
    document = ingest_excel_file(db_session, filename=filename, file_path=path)
    document.confidentiality_tag = tag
    for chunk in document.chunks:
        chunk.confidentiality_tag = tag
    db_session.commit()
    return document


@pytest.fixture
def cleanup(db_session):
    created = []
    yield created
    for document in created:
        db_session.delete(document)
    db_session.commit()


def test_roa_end_to_end_with_verified_figures_from_two_documents(monkeypatch, db_session, tmp_path, cleanup):
    income = _ingest(db_session, tmp_path, "EERR_Andina_2024_be010.xlsx",
                     [["Concepto", "2024"], ["Ingresos", 4820000], ["Utilidad neta", 610000]], sheet="Estado de resultados")
    balance = _ingest(db_session, tmp_path, "Balance_Andina_2024_be010.xlsx",
                      [["Cuenta", "2024"], ["Total activos", 6100000], ["Short-Term Obligations", 1200000]],
                      sheet="Balance general")
    cleanup.extend([income, balance])
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: _FakeExtractionModel({
        income.filename: {LINE_ITEMS["net_income"].query: ("Utilidad neta", "610,000")},
        balance.filename: {LINE_ITEMS["total_assets"].query: ("Total activos", "6,100,000")},
    }))

    report = RatioReportOut.model_validate(
        compute_ratios(RatioRequest(document_ids=[income.id, balance.id], ratios=["roa"]), db=db_session)
    )
    roa = report.ratios[0]
    assert roa.status == "calculated"
    assert roa.value == pytest.approx(0.1)
    assert roa.cross_document
    assert {c.document_id for c in roa.citations} == {str(income.id), str(balance.id)}
    assert any("610000" in c.text.replace(",", "") for c in roa.citations)


def test_a_figure_the_model_invents_is_not_used(monkeypatch, db_session, tmp_path, cleanup):
    income = _ingest(db_session, tmp_path, "EERR_Pacifico_2024_be010.xlsx",
                     [["Concepto", "2024"], ["Ingresos", 4820000], ["EBITDA", 1205000]])
    cleanup.append(income)
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: _FakeExtractionModel({
        income.filename: {
            LINE_ITEMS["revenue"].query: ("Ingresos", "4,820,000"),
            LINE_ITEMS["ebitda"].query: ("EBITDA", "9,999,999"),  # not in the document
        },
    }))

    report = compute_ratios(RatioRequest(document_ids=[income.id], ratios=["ebitda_margin"]), db=db_session)
    ebitda_margin = report.ratios[0]
    assert ebitda_margin.status == "not_calculable" and ebitda_margin.value is None
    assert "EBITDA" in ebitda_margin.reasons[0]


# --- review fixes: label doubts, coverage, signs, concurrency, model failures ---------------

def test_query_includes_common_synonyms():
    assert LINE_ITEMS["revenue"].query == "Revenue / Ingresos / Sales / Net sales"
    assert LINE_ITEMS["ebitda"].query == "EBITDA / Adjusted EBITDA / EBITDA ajustado"


def _found(document, label, value, **kwargs):
    return ExtractedFigure(document_id=str(document.id), found=True, value=float(value), value_text=f"{value:,}",
                           label=label, citation=_citation(document, label, value), **kwargs)


def test_model_doubt_about_an_accepted_label_is_dropped(statements):
    income, _ = statements
    figure = _found(income, "Net sales", 4_820_000, confidence="medium",
                    issues={"label": "The line item is labelled 'Net sales', which may not be exactly Revenue."},
                    reason="The line item is labelled 'Net sales', which may not be exactly Revenue.")
    mapped = to_mapped_input(LINE_ITEMS["revenue"], income, figure)
    assert mapped.confidence == "high" and mapped.reason is None and not mapped.needs_review


def test_other_extraction_issues_are_kept(statements):
    income, _ = statements
    figure = _found(income, "Net sales", 4_820_000, confidence="medium",
                    issues={"label": "label doubt", "periods": "The document also has FY2023."},
                    reason="label doubt The document also has FY2023.")
    mapped = to_mapped_input(LINE_ITEMS["revenue"], income, figure)
    assert mapped.confidence == "medium"
    assert mapped.reason == "The document also has FY2023."


def test_an_accepted_label_stops_the_search_at_the_first_document(statements):
    income, balance = statements
    tables = {income: {"revenue": ("Net sales", 4_820_000)}, balance: {"revenue": ("Revenue", 4_820_000)}}
    extractor = FakeExtractor(tables, overrides={"revenue": {"confidence": "medium", "issues": {"label": "doubt"}}})
    map_line_items(None, [income, balance], ["revenue"], extractor=extractor)
    assert len(extractor.calls) == 1


def test_quarter_and_year_to_date_inputs_are_held_for_review(statements):
    income, _ = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT}, [income], ratio_keys=["net_margin"], overrides={
        "net_income": {"period_coverage": "nine_months"}, "revenue": {"period_coverage": "three_months"},
    })
    assert ratios["net_margin"].status == "needs_review"
    assert "nine months" in ratios["net_margin"].reasons[0]


def test_balance_sheet_dates_are_not_a_coverage_mix(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"],
                        overrides={"net_income": {"period_coverage": "twelve_months"},
                                   "total_assets": {"period_coverage": "point_in_time"}})
    assert ratios["roa"].status == "calculated"


def test_negative_equity_is_flagged(statements):
    _, balance = statements
    table = {**BALANCE_SHEET, "total_equity": ("Total equity", -500_000)}
    income = statements[0]
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: table}, [income, balance], ratio_keys=["debt_to_equity", "roe"])
    for key in ("debt_to_equity", "roe"):
        assert ratios[key].status == "needs_review", key
        assert any("negative" in reason for reason in ratios[key].reasons), key


def test_negative_profit_is_not_flagged(statements):
    income, _ = statements
    table = {**INCOME_STATEMENT, "net_income": ("Net income", -120_000)}
    _, ratios, _ = _run({income: table}, [income], ratio_keys=["net_margin"])
    assert ratios["net_margin"].status == "calculated"
    assert ratios["net_margin"].value == pytest.approx(-120_000 / 4_820_000)


class _FakeSession:
    def __init__(self, log):
        self.log = log

    def commit(self):
        self.log.append("commit")

    def close(self):
        self.log.append("close")


def test_inputs_are_looked_up_concurrently_with_one_session_each(statements):
    income, balance = statements
    log = []
    extractor = FakeExtractor({income: INCOME_STATEMENT, balance: BALANCE_SHEET})
    report = calculate_ratios(None, [income, balance], extractor=extractor,
                              session_factory=lambda: _FakeSession(log), max_workers=4)
    ratios = {r.key: r for r in report.ratios}
    for key, expected in EXPECTED.items():
        assert ratios[key].value == pytest.approx(expected), key
    items = len(LINE_ITEMS)
    assert log.count("commit") == items and log.count("close") == items


class _FailingExtractor(FakeExtractor):
    def __init__(self, tables, failing):
        super().__init__(tables)
        self.failing = failing

    def __call__(self, db, document_id, metric, **kwargs):
        if self.failing is None or any(LINE_ITEMS[k].query == metric for k in self.failing):
            return ExtractedFigure(document_id=str(document_id), found=False, error=True,
                                   reason="The language model service didn't respond.")
        return super().__call__(db, document_id, metric, **kwargs)


def test_model_down_for_every_input_raises(statements):
    income, _ = statements
    with pytest.raises(ExtractionServiceError):
        calculate_ratios(None, [income], ratio_keys=["net_margin"],
                         extractor=_FailingExtractor({income: INCOME_STATEMENT}, failing=None))


def test_model_down_for_some_inputs_is_noted(statements):
    income, _ = statements
    report = calculate_ratios(None, [income], ratio_keys=["net_margin", "gross_margin"],
                              extractor=_FailingExtractor({income: INCOME_STATEMENT}, failing=["net_income"]))
    ratios = {r.key: r for r in report.ratios}
    assert ratios["net_margin"].status == "not_calculable"
    assert ratios["gross_margin"].status == "calculated"
    assert any("didn't respond for Net income" in note for note in report.notes)


class _BrokenModel:
    def with_structured_output(self, schema):
        class _Structured:
            def invoke(self, messages):
                raise ConnectionError("model service down")
        return _Structured()


def test_route_answers_503_when_the_model_is_down(monkeypatch, db_session, tmp_path, cleanup):
    income = _ingest(db_session, tmp_path, "EERR_down_be010.xlsx", [["Concepto", "2024"], ["Ingresos", 4820000]])
    cleanup.append(income)
    monkeypatch.setattr(extraction_module, "get_chat_model", lambda: _BrokenModel())
    with pytest.raises(HTTPException) as exc:
        compute_ratios(RatioRequest(document_ids=[income.id], ratios=["net_margin"]), db=db_session)
    assert exc.value.status_code == 503


def test_input_from_the_wrong_column_for_the_requested_period_is_held(statements):
    income, _ = statements
    coverage = "The value covers nine months (Nine Months Ended), which doesn't match the period requested (Q3 2026)."
    _, ratios, _ = _run({income: INCOME_STATEMENT}, [income], ratio_keys=["net_margin", "gross_margin"], period="Q3 2026",
                        overrides={"net_income": {"confidence": "medium", "issues": {"coverage": coverage}, "reason": coverage}})
    assert ratios["net_margin"].status == "needs_review"
    assert any("Nine Months Ended" in reason for reason in ratios["net_margin"].reasons)
    assert ratios["gross_margin"].status == "calculated"


@pytest.mark.parametrize(
    ("key", "label"),
    [("cost_of_sales", "Cost of revenues"), ("accounts_receivable", "Accounts receivable, net"),
     ("total_equity", "Total stockholders’ equity")],
)
def test_us_gaap_labels_are_standard(key, label):
    assert label_matches(LINE_ITEMS[key], label)


def test_balance_sheet_queries_ask_for_the_period_end_balance():
    assert LINE_ITEMS["inventory"].query.endswith("(balance sheet amount at the period-end date)")
    assert "balance sheet" not in LINE_ITEMS["net_income"].query


def test_balance_taken_from_a_cash_flow_column_is_held(statements):
    _, balance = statements
    figure = _found(balance, "Inventory", -6_696, confidence="high", period_coverage="three_months",
                    column_header="Quarter Ended June 30 2026")
    mapped = to_mapped_input(LINE_ITEMS["inventory"], balance, figure)
    assert mapped.coverage_mismatch and mapped.confidence == "medium"
    assert "cash-flow change" in mapped.reason


def test_same_date_written_differently_is_one_period(statements):
    _, balance = statements
    _, ratios, _ = _run({balance: BALANCE_SHEET}, [balance], ratio_keys=["current_ratio"],
                        periods={"current_assets": "As of June 30, 2026", "current_liabilities": "June 30, 2026"})
    assert ratios["current_ratio"].status == "calculated"


def test_quarter_end_balance_and_quarter_flows_are_one_period(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"],
                        periods={"net_income": "Q2 2026", "total_assets": "As of June 30, 2026"})
    assert ratios["roa"].status == "calculated"


def test_unknown_period_on_some_inputs_does_not_suggest_a_period(statements):
    income, balance = statements
    report, _, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance],
                        ratio_keys=["roa"], periods={"total_assets": None})
    assert not any("Specify a period" in note for note in report.notes)


def test_inputs_from_different_companies_are_not_combined(statements):
    # Andina's balance sheet with another company's (same currency) liabilities.
    income, balance = statements
    other = _document("Balance_Pacifico_2024.xlsx", ["Balance general"])
    tables = {balance: {"current_assets": ("Total current assets", 2_100_000)},
              other: {"current_liabilities": ("Pasivos corrientes", 1_600_000)}}
    report, ratios, _ = _run(tables, [balance, other], ratio_keys=["current_ratio"], overrides={
        "current_assets": {"entity": "Andina S.A.S."}, "current_liabilities": {"entity": "Pacífico S.A."},
    })
    assert ratios["current_ratio"].status == "not_calculable" and ratios["current_ratio"].value is None
    assert "different companies" in ratios["current_ratio"].reasons[0]
    assert "different companies" in report.notes[0]


def test_same_company_written_differently_is_combined(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"],
                        overrides={"net_income": {"entity": "Andina"}, "total_assets": {"entity": "ANDINA S.A.S."}})
    assert ratios["roa"].status == "calculated"


def test_an_undocumented_scale_claim_is_not_a_review_reason(statements):
    _, balance = statements
    claim = "The figures may be in thousands, but the document doesn't say so; the value is shown as written."
    figure = _found(balance, "Total current assets", 2_100_000, confidence="medium",
                    issues={"scale": claim}, reason=claim)
    mapped = to_mapped_input(LINE_ITEMS["current_assets"], balance, figure)
    assert mapped.confidence == "high" and mapped.reason is None


def test_inputs_are_compared_by_the_end_date_of_their_column(statements):
    income, balance = statements
    _, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance], ratio_keys=["roa"],
                        overrides={"net_income": {"period": "Q3 2026", "column_header": "Three Months Ended June 27, 2026"},
                                   "total_assets": {"period": "June 27, 2026", "column_header": "June 27, 2026"}})
    assert ratios["roa"].status == "calculated"


def test_section_prefixed_labels_still_map():
    assert label_matches(LINE_ITEMS["inventory"], "Current assets: Inventories")
    assert not label_matches(LINE_ITEMS["revenue"], "Cost of sales: Products")


def test_flow_read_from_unlabelled_multi_period_text_is_held(statements):
    income, _ = statements
    text = "STATEMENTS OF OPERATIONS Three Months Ended Nine Months Ended ... Total cost of sales 54,647 50,318 185,575 166,835"
    figure = _found(income, "Total cost of sales", 166_835, confidence="high")
    figure.citation = ResolvedCitation("S1", str(income.id), income.filename, None, None, 4, None, text)
    mapped = to_mapped_input(LINE_ITEMS["cost_of_sales"], income, figure)
    assert mapped.coverage_mismatch and "without column labels" in mapped.reason

    # Any PDF text without labels, even when the window lost the "Months Ended" header.
    plain = _found(income, "Total cost of sales", 166_835, confidence="high", period_coverage="three_months")
    plain.citation = ResolvedCitation("S1", str(income.id), income.filename, None, None, 4, None, "Total cost of sales 54,647 50,318 185,575 166,835")
    assert to_mapped_input(LINE_ITEMS["cost_of_sales"], income, plain).coverage_mismatch


def test_flow_from_a_labelled_row_is_not_held(statements):
    income, _ = statements
    text = "row 'Total cost of sales': Three Months Ended June 27, 2026=54,647; Nine Months Ended June 27, 2026=185,575"
    figure = _found(income, "Total cost of sales", 54_647, confidence="high", period_coverage="three_months")
    figure.citation = ResolvedCitation("S1", str(income.id), income.filename, None, None, 4, None, text)
    assert not to_mapped_input(LINE_ITEMS["cost_of_sales"], income, figure).coverage_mismatch


def test_a_flow_covering_another_span_than_the_rest_is_held(statements):
    income, balance = statements
    report, ratios, _ = _run({income: INCOME_STATEMENT, balance: BALANCE_SHEET}, [income, balance],
                             ratio_keys=["inventory_turnover", "gross_margin"], overrides={
                                 "cost_of_sales": {"period_coverage": "nine_months"},
                                 "revenue": {"period_coverage": "three_months"},
                                 "gross_profit": {"period_coverage": "three_months"}})
    assert ratios["inventory_turnover"].status == "needs_review"
    assert any("covers nine months" in reason for reason in ratios["inventory_turnover"].reasons)
    assert ratios["gross_margin"].status == "calculated"

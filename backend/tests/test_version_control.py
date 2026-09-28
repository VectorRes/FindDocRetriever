"""ID-HU-BE-015 — version supersession and default version selection.

Requires a reachable database (run via `docker compose up`), same as
test_document_resolution.py. Filenames get a random suffix so each test's
version group is isolated from other documents in the database.
"""
import uuid

import openpyxl
import pytest
from fastapi import HTTPException

from app.api.routes_documents import (
    approve_document,
    delete_document,
    list_versions,
    resolve_document,
)
from app.ingestion.service import ingest_excel_file
from app.qa.schemas import GroundedStatement, ResolvedCitation
from app.qa.service import describe_versions_used
from app.retrieval.service import retrieve
from app.versioning import derive_version_group

# fake_embeddings and db_session fixtures come from conftest.py


def _ingest(db_session, tmp_path, filename: str, marker: str = "1", version_of=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = f"zorblat budget total {marker}"
    ws["B1"] = int(marker) if marker.isdigit() else 0
    file_path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
    wb.save(file_path)
    return ingest_excel_file(db_session, filename=filename, file_path=file_path, version_of=version_of)


def _cleanup(db_session, *documents):
    for document in documents:
        db_session.delete(document)
    db_session.commit()


@pytest.fixture
def base_name():
    return f"Budget_{uuid.uuid4().hex[:8]}"


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("Budget_v1.xlsx", "excel:budget"),
        ("Budget_v2.xlsx", "excel:budget"),
        ("budget.xlsx", "excel:budget"),
        ("Budget v2.1.xlsx", "excel:budget"),
        ("Budget-ver3.xlsx", "excel:budget"),
        ("Budget (1).xlsx", "excel:budget"),
        ("Q2_CashFlow_v1.xlsx", "excel:q2 cashflow"),
        ("Budget_2025_v1.xlsx", "excel:budget 2025"),
        # "nov2" is not a version marker: no separator before the "v".
        ("Sales_nov2.xlsx", "excel:sales nov2"),
    ],
)
def test_derive_version_group(filename, expected):
    assert derive_version_group(filename, "excel") == expected


def test_pdf_and_excel_with_same_name_are_different_groups():
    assert derive_version_group("Budget_v1.pdf", "pdf") != derive_version_group("Budget_v1.xlsx", "excel")


def test_first_upload_is_the_default_even_before_approval(db_session, tmp_path, base_name):
    doc = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    try:
        assert doc.version_number == 1
        assert doc.approval_status == "draft"
        assert doc.is_current is True
    finally:
        _cleanup(db_session, doc)


def test_new_draft_does_not_replace_the_approved_version_until_approved(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    approve_document(str(v1.id), x_user_role="reviewer", db=db_session)
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx")
    try:
        db_session.refresh(v1)
        assert v2.version_group == v1.version_group
        assert v2.version_number == 2
        # v2 is pending approval: v1 stays the default and is not superseded.
        assert v1.is_current is True and v1.superseded_by_id is None
        assert v2.is_current is False and v2.superseded_by_id is None

        approve_document(str(v2.id), x_user_role="reviewer", db=db_session)
        db_session.refresh(v1)
        db_session.refresh(v2)
        assert v2.is_current is True
        assert v1.is_current is False
        assert v1.superseded_by_id == v2.id
    finally:
        _cleanup(db_session, v1, v2)


def test_reuploading_without_approvals_makes_the_latest_version_the_default(db_session, tmp_path, base_name):
    """Keeps FE-005's behavior when nothing in the group has been approved:
    re-uploading the same filename supersedes the previous version, and a
    chain of versions links each one to its successor."""
    v1 = _ingest(db_session, tmp_path, f"{base_name}.xlsx")
    v2 = _ingest(db_session, tmp_path, f"{base_name}.xlsx")
    v3 = _ingest(db_session, tmp_path, f"{base_name}.xlsx")
    try:
        for doc in (v1, v2, v3):
            db_session.refresh(doc)
        assert [d.version_number for d in (v1, v2, v3)] == [1, 2, 3]
        assert v3.is_current is True
        assert v1.superseded_by_id == v2.id
        assert v2.superseded_by_id == v3.id
    finally:
        _cleanup(db_session, v1, v2, v3)


def test_approve_requires_a_reviewer_role(db_session, tmp_path, base_name):
    doc = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    try:
        with pytest.raises(HTTPException) as exc_info:
            approve_document(str(doc.id), x_user_role="analyst", db=db_session)
        assert exc_info.value.status_code == 403
        db_session.refresh(doc)
        assert doc.approval_status == "draft"
    finally:
        _cleanup(db_session, doc)


def test_retrieval_uses_the_default_version_unless_scoped_to_another(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx", marker="100")
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx", marker="200")
    ours = {str(v1.id), str(v2.id)}
    try:
        default_results = [
            r for r in retrieve(db_session, "zorblat budget total", top_k=20, user_roles="admin")
            if r.document_id in ours
        ]
        assert default_results
        assert {r.document_id for r in default_results} == {str(v2.id)}

        audit_results = retrieve(
            db_session, "zorblat budget total", top_k=20, user_roles="admin", document_id=v1.id
        )
        assert audit_results
        assert {r.document_id for r in audit_results} == {str(v1.id)}
    finally:
        db_session.rollback()
        _cleanup(db_session, v1, v2)


def test_deleting_the_default_version_falls_back_to_the_previous_one(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx")
    try:
        delete_document(str(v2.id), db=db_session)
        db_session.refresh(v1)
        assert v1.is_current is True
        assert v1.superseded_by_id is None
    finally:
        _cleanup(db_session, v1)


def test_superseded_versions_stay_listed_and_resolvable_for_audit(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx")
    try:
        versions = list_versions(str(v1.id), db=db_session).documents
        assert [d.id for d in versions] == [v2.id, v1.id]

        resolution = resolve_document(str(v1.id), x_user_role="admin", db=db_session)
        assert resolution.is_superseded is True
        assert resolution.current_version.id == v2.id
    finally:
        _cleanup(db_session, v1, v2)


def test_pending_draft_is_not_reported_as_superseded(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    approve_document(str(v1.id), x_user_role="admin", db=db_session)
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx")
    try:
        resolution = resolve_document(str(v2.id), x_user_role="admin", db=db_session)
        assert resolution.is_superseded is False
        assert resolution.document.approval_status == "draft"
    finally:
        _cleanup(db_session, v1, v2)


def test_version_of_attaches_an_unrelated_filename_to_a_group(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    renamed = _ingest(db_session, tmp_path, f"Presupuesto_final_{uuid.uuid4().hex[:6]}.xlsx", version_of=v1)
    try:
        db_session.refresh(v1)
        assert renamed.version_group == v1.version_group
        assert renamed.version_number == 2
        assert v1.superseded_by_id == renamed.id
    finally:
        _cleanup(db_session, v1, renamed)


def test_answer_reports_which_version_was_used(db_session, tmp_path, base_name):
    v1 = _ingest(db_session, tmp_path, f"{base_name}_v1.xlsx")
    v2 = _ingest(db_session, tmp_path, f"{base_name}_v2.xlsx")

    def _statement(doc):
        citation = ResolvedCitation(
            source_id="S1", document_id=str(doc.id), document_filename=doc.filename,
            sheet_name="Budget", cell_range="B1", page_number=None, reference_number=None, text="",
        )
        return GroundedStatement(text="Total was X.", citations=[citation, citation])

    try:
        [current] = describe_versions_used(db_session, [_statement(v2)])
        assert current.filename == v2.filename
        assert current.version_number == 2
        assert current.is_current is True
        assert current.current_version_filename is None

        [audited] = describe_versions_used(db_session, [_statement(v1)])
        assert audited.is_current is False
        assert audited.current_version_filename == v2.filename
    finally:
        _cleanup(db_session, v1, v2)

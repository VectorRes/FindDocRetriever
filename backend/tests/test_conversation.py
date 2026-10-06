"""Integration tests for ID-HU-FE-001's follow-up conversation context.

Requires a reachable database (run via `docker compose up`), same as
test_retrieval.py — skips otherwise. The chat model is monkeypatched, same
approach as test_qa_generation.py, but here we also assert on *what prompt*
was sent, since the behavior under test is "does history reach the model".
"""
from uuid import UUID

import openpyxl

import app.qa.graph as qa_graph_module
from app.db.models import ConversationSession
from app.ingestion.service import ingest_excel_file
from app.qa.schemas import LLMCitation, LLMGroundedAnswer, LLMStatement, LLMTriage
from app.qa.service import answer_question

# fake_embeddings and db_session fixtures come from conftest.py


class _RecordingStructuredLLM:
    def __init__(self, response: LLMGroundedAnswer, prompts: list[str]):
        self._response = response
        self._prompts = prompts

    def invoke(self, messages):
        self._prompts.append(messages[-1]["content"])
        return self._response


class _RecordingChatModel:
    def __init__(self, response: LLMGroundedAnswer, prompts: list[str]):
        self._response = response
        self._prompts = prompts

    def with_structured_output(self, schema):
        if schema is LLMTriage:
            # ID-HU-BE-009's pre-check: not ambiguous, and not recorded, so
            # `prompts` holds only the answer-generation prompts.
            return _RecordingStructuredLLM(LLMTriage(), [])
        return _RecordingStructuredLLM(self._response, self._prompts)


def _patch_llm(monkeypatch, response: LLMGroundedAnswer, prompts: list[str]) -> None:
    monkeypatch.setattr(qa_graph_module, "get_chat_model", lambda: _RecordingChatModel(response, prompts))


def test_follow_up_question_receives_prior_turn_as_history(monkeypatch, tmp_path, db_session):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Line Item"
    ws["B1"] = "Q1"
    ws["A2"] = "Revenue"
    ws["B2"] = 120000
    file_path = tmp_path / "budget.xlsx"
    wb.save(file_path)

    document = ingest_excel_file(db_session, filename="budget.xlsx", file_path=file_path)
    document.confidentiality_tag = "public"
    for chunk in document.chunks:
        chunk.confidentiality_tag = "public"
    db_session.commit()
    session_id = None
    try:
        prompts: list[str] = []
        raw = LLMGroundedAnswer(
            statements=[
                # Cites both retrieved chunks: their order isn't guaranteed, and
                # BE-009's figure check needs the cited sources to contain 120,000.
                LLMStatement(
                    text="Revenue was 120,000.",
                    citations=[LLMCitation(source_id="S1"), LLMCitation(source_id="S2")],
                )
            ]
        )
        _patch_llm(monkeypatch, raw, prompts)

        first = answer_question(db_session, question="What was Revenue?")
        assert first.session_id is not None
        assert "CONVERSATION HISTORY" not in prompts[0]

        second = answer_question(
            db_session, question="And what about last quarter?", session_id=UUID(first.session_id)
        )
        session_id = second.session_id

        assert second.session_id == first.session_id
        assert "CONVERSATION HISTORY" in prompts[1]
        assert "What was Revenue?" in prompts[1]
        assert "Revenue was 120,000." in prompts[1]
    finally:
        db_session.delete(document)
        if session_id is not None:
            session = db_session.get(ConversationSession, UUID(session_id))
            if session is not None:
                db_session.delete(session)
        db_session.commit()

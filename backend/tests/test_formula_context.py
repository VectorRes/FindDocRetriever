"""Formula precedents added as QA sources (app/qa/formula_context.py)."""
import uuid

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from app.db.models import Chunk, RetrievalAuditLog
from app.ingestion.service import ingest_excel_file
from app.qa.formula_context import MAX_PRECEDENTS, add_formula_precedents
from app.retrieval.service import RetrievedChunk

# db_session and fake_embeddings fixtures come from conftest.py


def _budget(db_session, tmp_path, tag="public"):
    wb = openpyxl.Workbook()
    assumptions = wb.active
    assumptions.title = "Supuestos"
    for row in [("Supuesto", "Valor"), ("Crecimiento", 0.10), ("Margen bruto", 0.41),
                ("Gastos (% ingresos)", 0.20), ("Tasa de impuesto", 0.35)]:
        assumptions.append(row)
    wb.defined_names["TasaImpuesto"] = DefinedName("TasaImpuesto", attr_text="Supuestos!$B$5")
    projection = wb.create_sheet("Proyección")
    for row in [("Concepto", "2025"), ("Ingresos", "=4820000*(1+Supuestos!B2)"), ("Utilidad bruta", "=B2*Supuestos!B3"),
                ("Gastos", "=B2*Supuestos!B4"), ("Utilidad operacional", "=B3-B4"),
                ("Impuesto", "=B5*TasaImpuesto"), ("Utilidad neta", "=B5-B6")]:
        projection.append(row)
    path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
    wb.save(path)
    document = ingest_excel_file(db_session, filename=f"Presupuesto_{uuid.uuid4().hex[:6]}.xlsx", file_path=path)
    document.confidentiality_tag = tag
    for chunk in document.chunks:
        chunk.confidentiality_tag = tag
    db_session.commit()
    return document


@pytest.fixture
def cleanup(db_session):
    created = []
    yield created
    for document in created:
        db_session.delete(document)
    db_session.commit()


def _retrieved(document, sheet, cell_range) -> RetrievedChunk:
    chunk = next(c for c in document.chunks if c.sheet_name == sheet and c.cell_range == cell_range)
    return RetrievedChunk(
        chunk_id=str(chunk.id), document_id=str(document.id), document_filename=document.filename,
        sheet_name=chunk.sheet_name, cell_range=chunk.cell_range, page_number=None, text=chunk.text,
        confidentiality_tag=chunk.confidentiality_tag,
    )


def test_rows_a_formula_references_are_added(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _retrieved(document, "Proyección", "A7:B7")  # =B5-B6

    sources = add_formula_precedents(db_session, [net_income], "¿Cómo se calcula la utilidad neta?")

    added = {(s.sheet_name, s.cell_range) for s in sources[1:]}
    assert added == {("Proyección", "A5:B5"), ("Proyección", "A6:B6")}
    assert sources[0] is net_income


def test_named_range_is_resolved_and_explained(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    tax = _retrieved(document, "Proyección", "A6:B6")  # =B5*TasaImpuesto

    sources = add_formula_precedents(db_session, [tax], "¿Qué tasa de impuesto se usa?")

    rate = next(s for s in sources if s.sheet_name == "Supuestos")
    assert rate.cell_range == "A5:B5"
    assert rate.text.startswith("Named range TasaImpuesto = Supuestos!$B$5.")
    assert "0.35" in rate.text


def test_sources_without_formulas_are_left_alone(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    assumption = _retrieved(document, "Supuestos", "A5:B5")

    assert add_formula_precedents(db_session, [assumption], "tasa") == [assumption]


def test_already_retrieved_rows_are_not_duplicated(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _retrieved(document, "Proyección", "A7:B7")
    operating = _retrieved(document, "Proyección", "A5:B5")

    sources = add_formula_precedents(db_session, [net_income, operating], "utilidad neta")

    ranges = [(s.sheet_name, s.cell_range) for s in sources]
    assert len(ranges) == len(set(ranges))


def test_restricted_precedents_are_not_added(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    for chunk in document.chunks:
        if chunk.sheet_name == "Supuestos":
            chunk.confidentiality_tag = "restricted"
    db_session.commit()
    tax = _retrieved(document, "Proyección", "A6:B6")

    sources = add_formula_precedents(db_session, [tax], "impuesto", user_roles="analyst")

    assert not any(s.sheet_name == "Supuestos" for s in sources)


def test_added_precedents_are_audited_and_capped(db_session, tmp_path, cleanup):
    document = _budget(db_session, tmp_path)
    cleanup.append(document)
    every_formula_row = [_retrieved(document, "Proyección", f"A{row}:B{row}") for row in range(2, 8)]

    sources = add_formula_precedents(db_session, every_formula_row, "presupuesto", user_id="ana")

    added = sources[len(every_formula_row):]
    assert 0 < len(added) <= MAX_PRECEDENTS
    audited = db_session.query(RetrievalAuditLog).filter(
        RetrievalAuditLog.user_id == "ana", RetrievalAuditLog.document_id == document.id
    ).count()
    assert audited == len(added)
    assert all(db_session.get(Chunk, uuid.UUID(s.chunk_id)) is not None for s in added)


# --- explain_formulas: the calculation spelled out from the workbook ---------------------

from app.qa.formula_context import asks_for_calculation, explain_formulas  # noqa: E402
from app.qa.schemas import GroundedStatement, ResolvedCitation  # noqa: E402


def _statement_budget(db_session, tmp_path):
    wb = openpyxl.Workbook()
    assumptions = wb.active
    assumptions.title = "Supuestos"
    assumptions.append(("Supuesto", "Valor"))
    assumptions.append(("Tasa de impuesto", 0.35))
    wb.defined_names["TasaImpuesto"] = DefinedName("TasaImpuesto", attr_text="Supuestos!$B$2")
    projection = wb.create_sheet("Proyección")
    for row in [("Concepto", "2025"), ("Utilidad operacional", 1113420), ("Impuesto", "=B2*TasaImpuesto"),
                ("Utilidad neta", "=B2-B3")]:
        projection.append(row)
    path = tmp_path / f"{uuid.uuid4().hex}.xlsx"
    wb.save(path)
    document = ingest_excel_file(db_session, filename=f"Presupuesto_{uuid.uuid4().hex[:6]}.xlsx", file_path=path)
    for chunk in document.chunks:
        chunk.confidentiality_tag = "public"
    db_session.commit()
    return document


def _citation(document, sheet, cell_range) -> ResolvedCitation:
    chunk = next(c for c in document.chunks if c.sheet_name == sheet and c.cell_range == cell_range)
    return ResolvedCitation("S1", str(document.id), document.filename, sheet, cell_range, None, None, chunk.text)


@pytest.mark.parametrize(
    ("question", "expected"),
    [("¿Cómo se calcula la utilidad neta presupuestada?", True), ("¿De dónde sale el impuesto?", True),
     ("How is budgeted net income calculated?", True), ("What formula gives the tax?", True),
     ("¿Cuál fue la utilidad neta en 2024?", False), ("What was revenue?", False)],
)
def test_asks_for_calculation(question, expected):
    assert asks_for_calculation(question) is expected


def test_formula_is_explained_with_line_items_and_values(db_session, tmp_path, cleanup):
    document = _statement_budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _citation(document, "Proyección", "A4:B4")
    sources = [net_income, _citation(document, "Proyección", "A2:B2"), _citation(document, "Proyección", "A3:B3")]
    answer = [GroundedStatement(text="La utilidad neta presupuestada es ...", citations=[net_income])]

    [explanation] = explain_formulas(db_session, "¿Cómo se calcula la utilidad neta?", answer, sources)

    assert explanation.text.startswith("Cálculo de Utilidad neta")
    assert "=B2-B3" in explanation.text
    assert "B2 = Utilidad operacional (1,113,420)" in explanation.text
    assert "B3 = Impuesto" in explanation.text
    assert [c.cell_range for c in explanation.citations] == ["A4:B4", "A2:B2", "A3:B3"]


def test_named_range_component_is_named_and_cited(db_session, tmp_path, cleanup):
    document = _statement_budget(db_session, tmp_path)
    cleanup.append(document)
    tax = _citation(document, "Proyección", "A3:B3")
    rate = _citation(document, "Supuestos", "A2:B2")
    operating = _citation(document, "Proyección", "A2:B2")
    answer = [GroundedStatement(text="The budgeted tax is ...", citations=[tax])]

    [explanation] = explain_formulas(db_session, "How is the tax calculated?", answer, [tax, operating, rate])

    assert explanation.text.startswith("How Impuesto")
    assert "TasaImpuesto = Tasa de impuesto (0.35)" in explanation.text
    assert rate in explanation.citations


def test_no_explanation_when_not_asked_or_already_given(db_session, tmp_path, cleanup):
    document = _statement_budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _citation(document, "Proyección", "A4:B4")
    operating = _citation(document, "Proyección", "A2:B2")
    sources = [net_income, operating, _citation(document, "Proyección", "A3:B3")]
    answer = [GroundedStatement(text="...", citations=[net_income])]

    assert explain_formulas(db_session, "¿Cuál es la utilidad neta presupuestada?", answer, sources) == []
    already = answer + [GroundedStatement(text="Operating income is 1,113,420", citations=[operating])]
    assert explain_formulas(db_session, "¿Cómo se calcula la utilidad neta?", already, sources) == []


def test_components_without_a_given_source_are_not_described(db_session, tmp_path, cleanup):
    document = _statement_budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _citation(document, "Proyección", "A4:B4")
    answer = [GroundedStatement(text="...", citations=[net_income])]

    assert explain_formulas(db_session, "¿Cómo se calcula la utilidad neta?", answer, [net_income]) == []


def test_only_the_top_formula_is_considered_and_not_when_already_explained(db_session, tmp_path, cleanup):
    # The model explained net income itself (citing it and its tax input): the
    # tax formula is an input of net income, so it isn't explained on its own.
    document = _statement_budget(db_session, tmp_path)
    cleanup.append(document)
    net_income = _citation(document, "Proyección", "A4:B4")
    tax = _citation(document, "Proyección", "A3:B3")
    operating = _citation(document, "Proyección", "A2:B2")
    rate = _citation(document, "Supuestos", "A2:B2")
    answer = [GroundedStatement(text="Net income is operating income minus tax", citations=[net_income]),
              GroundedStatement(text="Tax is ...", citations=[tax])]

    assert explain_formulas(db_session, "How is net income calculated?", answer, [net_income, tax, operating, rate]) == []

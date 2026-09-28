"""Integration test: ingest a workbook into Postgres/pgvector, then query it.

Requires a reachable database (run via `docker compose up`, or point
DATABASE_URL at a local Postgres+pgvector instance). Skips otherwise.
"""
import openpyxl

from app.ingestion.service import ingest_excel_file
from app.retrieval.service import retrieve

# fake_embeddings and db_session fixtures come from conftest.py


def test_ingest_and_query_round_trip(tmp_path, db_session):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Line Item"
    ws["B1"] = "Q1"
    ws["A2"] = "Revenue"
    ws["B2"] = 120000
    ws["A3"] = "COGS"
    ws["B3"] = 45000
    file_path = tmp_path / "budget.xlsx"
    wb.save(file_path)

    document = ingest_excel_file(db_session, filename="budget.xlsx", file_path=file_path)
    try:
        assert document.status == "ready"
        document.confidentiality_tag = "public"
        for chunk in document.chunks:
            chunk.confidentiality_tag = "public"
        db_session.commit()

        results = retrieve(db_session, question="What is the Revenue line item?", top_k=5)

        assert results, "expected at least one retrieved chunk"
        assert any("Revenue" in r.text for r in results)
        assert results[0].document_filename == "budget.xlsx"
        top_hit = next(r for r in results if "Revenue" in r.text)
        assert top_hit.sheet_name == "Budget"
        assert top_hit.cell_range == "A2:B2"
    finally:
        # Ingestion commits directly (it isn't wrapped in the test's rollback-only
        # session), so without this, ingested rows pile up in Postgres across
        # runs and can contaminate other integration tests' top-k results.
        db_session.delete(document)
        db_session.commit()

def test_restricted_chunk_is_excluded_for_analyst_and_available_with_permission(db_session):
    from app.db.models import Chunk, Document

    public_doc = Document(
        filename="public.xlsx", doc_type="excel", status="ready",
        warnings=[], confidentiality_tag="public",
    )
    restricted_doc = Document(
        filename="compensation.xlsx", doc_type="excel", status="ready",
        warnings=[], confidentiality_tag="restricted:compensation",
    )
    db_session.add_all([public_doc, restricted_doc])
    db_session.flush()

    from app.embeddings.factory import get_embedding_provider
    provider = get_embedding_provider()
    texts = ["Revenue was 100000", "Employee compensation was 50000"]
    embeddings = provider.embed(texts)
    db_session.add_all([
        Chunk(document_id=public_doc.id, text=texts[0], embedding=embeddings[0],
              confidentiality_tag="public"),
        Chunk(document_id=restricted_doc.id, text=texts[1], embedding=embeddings[1],
              confidentiality_tag="restricted:compensation"),
    ])
    db_session.commit()

    try:
        analyst_results = retrieve(
            db_session, "What was employee compensation?", top_k=5, user_roles="analyst"
        )
        assert all("compensation" not in r.text.lower() for r in analyst_results)

        cleared_results = retrieve(
            db_session, "What was employee compensation?", top_k=5,
            user_roles="compensation-access",
        )
        assert any("compensation" in r.text.lower() for r in cleared_results)
    finally:
        db_session.delete(public_doc)
        db_session.delete(restricted_doc)
        db_session.commit()
